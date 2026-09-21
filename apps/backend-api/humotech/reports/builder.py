"""Конструктор отчётов: спецификация, область, строки, предпросмотр и файл.

Страница «Отчёты» выбирает вид, период, офисы, отделы, человека и поля.
Здесь это превращается в таблицу — одну и ту же для предпросмотра и
для файла. Два построителя «похожих» строк разошлись бы на первой же
правке, и предпросмотр перестал бы обещать то, что окажется в файле.

Правила, которые держатся здесь, а не на странице.

**Сутки — в поясе своего офиса.** Посещаемость собирается по офису за
раз, и каждый офис считается в своём часовом поясе. Сводка называет
пояса всех офисов выборки.

**Выходной и день без графика — не прогул.** Статус дня берётся у
`AttendanceHrService.presence()`, а он различает «не пришёл в рабочий
день», «выходной» и «графика нет».

**Отсутствия — только подтверждённые.** Отпуск и больничный в днях
посещаемости — это `EmployeeAbsence`, а она появляется лишь по
одобренной заявке. В отчёте «Отсутствия» видны все заявки со статусом,
но итоги по дням считаются только по одобренным.

**Незакрытая сессия — отдельно.** Её время не входит в «время в офисе»
и в фактические часы: конца у неё нет, и подставлять «сейчас» значило бы
выдумать выход.

**Будущих дней в отчёте нет.** «Не пришёл» про завтра — неправда.

**Область видимости не обходится.** Офисы проверяются по одному, строки
собирают те же сервисы, что и экраны.
"""

from __future__ import annotations

import csv
import io
import re
import time as monotonic
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from types import SimpleNamespace
from typing import Callable, Iterator
from zoneinfo import ZoneInfo

from django.db.models import Q

from humotech.core.errors import ValidationFailed
from humotech.core.rbac import Actor
from humotech.core.service import BaseService
from humotech.core.timeframes import days_in, office_zone, range_bounds, today
from humotech.reports.catalog import (
    DAY_KINDS, DAY_STATUSES, EVENT_SOURCES, EVENT_TYPES, FIELDS,
    KIND_PERMISSIONS, KIND_TITLES, REPORT_KINDS, REQUEST_STATUSES, STATUS_TITLES,
    TELEGRAM, VERIFICATION, Column, columns_for, default_fields,
)
from humotech.reports.export import XLSX_MAX_ROWS, Echo, safe_text
from humotech.reports.sheets import check_period

#: Сколько строк отдаёт предпросмотр.
PREVIEW_MIN_ROWS = 5
PREVIEW_MAX_ROWS = 20
#: Сколько дней периода просматривает предпросмотр. Этого хватает, чтобы
#: набрать строки и оценить частоту опозданий, и не хватает, чтобы
#: предпросмотр за год стал сборкой файла за год.
PREVIEW_DAYS = 7
#: Сколько офисов и отделов можно выбрать за раз.
MAX_CHOICES = 200
NAME_MAX = 120

#: Версия заказа. Старые задания без неё собирает `sheets.build_sheet`.
BUILDER_VERSION = 2

PERIOD_MODES = ("custom", "this_month", "last_month")
ABSENCE_STATES = frozenset({"SICK_LEAVE", "VACATION", "OTHER_ABSENCE"})


# --- спецификация -------------------------------------------------------------


@dataclass(frozen=True)
class ReportSpec:
    kind: str
    first: date
    last: date
    region_id: uuid.UUID | None = None
    office_ids: tuple[uuid.UUID, ...] = ()
    department_ids: tuple[uuid.UUID, ...] = ()
    employee_id: uuid.UUID | None = None
    include_inactive: bool = False
    fields: tuple[str, ...] = ()
    name: str | None = None
    period: str = "custom"

    @classmethod
    def build(
        cls,
        kind: str,
        *,
        date_from: date,
        date_to: date,
        region_id=None,
        office_ids=(),
        department_ids=(),
        employee_id=None,
        include_inactive: bool = False,
        fields=None,
        name: str | None = None,
        period: str | None = None,
    ) -> ReportSpec:
        if kind not in REPORT_KINDS:
            raise ValidationFailed(
                "Неизвестный вид отчёта",
                details={"kind": ["Такого отчёта нет"], "allowed": list(REPORT_KINDS)},
            )
        check_period(date_from, date_to)

        allowed = {item.key for item in FIELDS[kind]}
        chosen = default_fields(kind) if fields is None else tuple(dict.fromkeys(fields))
        unknown = [key for key in chosen if key not in allowed]
        if unknown:
            raise ValidationFailed(
                "Неизвестные поля отчёта",
                details={"fields": [f"Нет такого поля: {', '.join(unknown)}"]},
            )
        if not chosen:
            raise ValidationFailed(
                "Выберите хотя бы одно поле",
                details={"fields": ["Выберите хотя бы одно поле"]},
            )
        for label, values in (("office_ids", office_ids),
                              ("department_ids", department_ids)):
            if len(values) > MAX_CHOICES:
                raise ValidationFailed(
                    f"Можно выбрать не больше {MAX_CHOICES} значений",
                    details={label: [f"Не больше {MAX_CHOICES}"]},
                )
        return cls(
            kind=kind,
            first=date_from,
            last=date_to,
            region_id=_uuid(region_id),
            office_ids=tuple(dict.fromkeys(_uuid(v) for v in office_ids or ())),
            department_ids=tuple(dict.fromkeys(_uuid(v) for v in department_ids or ())),
            employee_id=_uuid(employee_id),
            include_inactive=bool(include_inactive),
            fields=chosen,
            name=clean_name(name),
            period=period if period in PERIOD_MODES else "custom",
        )

    def to_filters(self) -> dict:
        """Спецификация в JSON заказа. Её же открывает «Открыть параметры»."""
        return {
            "builder": BUILDER_VERSION,
            "date_from": self.first.isoformat(),
            "date_to": self.last.isoformat(),
            "region_id": str(self.region_id) if self.region_id else None,
            "office_ids": [str(v) for v in self.office_ids],
            "department_ids": [str(v) for v in self.department_ids],
            "employee_id": str(self.employee_id) if self.employee_id else None,
            "include_inactive": self.include_inactive,
            "fields": list(self.fields),
            "name": self.name,
            "period": self.period,
        }

    @classmethod
    def from_filters(cls, kind: str, filters: dict) -> ReportSpec:
        return cls.build(
            kind,
            date_from=date.fromisoformat(filters["date_from"]),
            date_to=date.fromisoformat(filters["date_to"]),
            region_id=filters.get("region_id"),
            office_ids=filters.get("office_ids") or (),
            department_ids=filters.get("department_ids") or (),
            employee_id=filters.get("employee_id"),
            include_inactive=filters.get("include_inactive", False),
            fields=filters.get("fields"),
            name=filters.get("name"),
            period=filters.get("period"),
        )

    def file_name(self, fmt: str) -> str:
        return f"{self.name or default_name(self.kind, self.first, self.last)}.{fmt}"


def is_builder_order(filters: dict | None) -> bool:
    return bool(filters) and filters.get("builder") == BUILDER_VERSION


def default_name(kind: str, first: date, last: date) -> str:
    """«Посещаемость_15.08–13.09.2026»: вид и период, без лишнего года."""
    if kind == "employees" or first == last:
        span = last.strftime("%d.%m.%Y")
    elif first.year == last.year:
        span = f"{first:%d.%m}–{last:%d.%m.%Y}"
    else:
        span = f"{first:%d.%m.%Y}–{last:%d.%m.%Y}"
    return f"{KIND_TITLES[kind]}_{span}"


_BAD_NAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


def clean_name(raw) -> str | None:
    """Имя файла от человека — без путей, управляющих символов и расширения."""
    if raw is None:
        return None
    text = _BAD_NAME.sub(" ", str(raw))
    text = re.sub(r"\.(xlsx|csv)\s*$", "", text.strip(), flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip(" .")
    return text[:NAME_MAX].strip() or None


def _uuid(value) -> uuid.UUID | None:
    if value in (None, ""):
        return None
    if isinstance(value, uuid.UUID):
        return value
    try:
        return uuid.UUID(str(value))
    except ValueError as exc:
        raise ValidationFailed(
            "Неверный идентификатор", details={"value": str(value)}
        ) from exc


# --- прогресс -----------------------------------------------------------------


class Progress:
    """Сколько сделано из скольких. Пишет в базу не чаще раза в секунду."""

    def __init__(self, sink: Callable[[int, int], None] | None = None) -> None:
        self.sink = sink
        self.done = 0
        self.rows = 0
        self._flushed = 0.0

    def step(self, count: int = 1) -> None:
        self.done += count
        self._flush()

    def row(self) -> None:
        self.rows += 1

    def _flush(self, force: bool = False) -> None:
        if self.sink is None:
            return
        now = monotonic.monotonic()
        if force or now - self._flushed >= 1.0:
            self._flushed = now
            self.sink(self.done, self.rows)

    def finish(self) -> None:
        self._flush(force=True)


# --- область ------------------------------------------------------------------


class Scope:
    """Офисы и сотрудники под фильтрами — внутри области видимости."""

    def __init__(self, service: BaseService, actor: Actor, spec: ReportSpec) -> None:
        from humotech.departments.models import Department
        from humotech.employees.selectors import require_visible_employee
        from humotech.offices.models import Office

        access = service.access
        self.actor = actor
        self.spec = spec

        offices = Office.objects.filter(organization_id=actor.organization_id)
        visible = access.office_filter(actor)
        if visible is not None:
            offices = offices.filter(visible)
        if spec.region_id:
            access.require_region(actor, spec.region_id)
            offices = offices.filter(region_id=spec.region_id)
        if spec.office_ids:
            for office_id in spec.office_ids:
                access.require_office(actor, office_id)
            offices = offices.filter(id__in=spec.office_ids)
        self.offices = list(offices.select_related("region").order_by("name"))

        self.employee = (
            require_visible_employee(access, actor, spec.employee_id)
            if spec.employee_id else None
        )
        self.departments = list(
            Department.objects.filter(
                organization_id=actor.organization_id, id__in=spec.department_ids
            ).order_by("name")
        ) if spec.department_ids else []

        self._assignments: dict | None = None

    @property
    def zones(self) -> dict[uuid.UUID, ZoneInfo]:
        return {office.id: office_zone(office) for office in self.offices}

    def assignments(self) -> dict:
        """Назначение каждого сотрудника выборки — последнее в периоде.

        Без неактивных — только работающие, чьё назначение действует на
        конец периода. С неактивными — все, чьё назначение пересекается
        с периодом: уволенный в середине месяца в отчёт за месяц попадает.
        """
        if self._assignments is not None:
            return self._assignments
        from humotech.employees.models import EmployeeAssignment
        from humotech.employees.services import WORKING_STATUSES

        spec = self.spec
        start = spec.first if spec.include_inactive else spec.last
        rows = EmployeeAssignment.objects.filter(
            is_primary=True,
            office_id__in=[office.id for office in self.offices],
            employee__organization_id=self.actor.organization_id,
            valid_from__lte=spec.last,
        ).filter(Q(valid_to__isnull=True) | Q(valid_to__gte=start))
        if not spec.include_inactive:
            # Стажёр работает и попадает в отчёт: «не включать неактивных»
            # означает уволенных, а не тех, у кого идёт испытательный срок.
            rows = rows.filter(employee__employment_status__in=WORKING_STATUSES)
        if spec.department_ids:
            rows = rows.filter(department_id__in=spec.department_ids)
        if spec.employee_id:
            rows = rows.filter(employee_id=spec.employee_id)
        rows = rows.select_related(
            "employee", "office", "office__region", "department", "position",
            "manager_employee",
        ).order_by("employee_id", "-valid_from")

        result: dict = {}
        for row in rows:
            result.setdefault(row.employee_id, row)
        self._assignments = result
        return result

    def past_days(self) -> list[date]:
        """Дни периода, которые уже наступили хотя бы в одном офисе."""
        if not self.offices:
            return []
        latest = max(today(zone) for zone in self.zones.values())
        return [day for day in days_in(self.spec.first, self.spec.last) if day <= latest]


# --- графики ------------------------------------------------------------------


class Schedules:
    """Действующие графики сотрудников за период. Один запрос на пачку людей."""

    def __init__(self, first: date, last: date) -> None:
        self.first = first
        self.last = last
        self._rows: dict[uuid.UUID, list] = {}

    def load(self, employee_ids) -> None:
        from humotech.schedules.models import EmployeeScheduleAssignment

        missing = [key for key in employee_ids if key not in self._rows]
        if not missing:
            return
        for key in missing:
            self._rows[key] = []
        rows = (
            EmployeeScheduleAssignment.objects.filter(
                employee_id__in=missing, valid_from__lte=self.last
            )
            .filter(Q(valid_to__isnull=True) | Q(valid_to__gte=self.first))
            .select_related("schedule")
            .prefetch_related("schedule__days__breaks")
            .order_by("employee_id", "-valid_from")
        )
        for row in rows:
            self._rows[row.employee_id].append(row)

    def on(self, employee_id: uuid.UUID, day: date):
        """График на день — то же правило, что у `presence()`: самый поздний."""
        for row in self._rows.get(employee_id, ()):
            if row.valid_from <= day and (row.valid_to is None or row.valid_to >= day):
                return row.schedule
        return None


def planned_seconds(row, schedule, day: date) -> int:
    """Плановое время дня: смена по графику минус неоплачиваемые перерывы."""
    if schedule is None or row.scheduled_start is None or row.scheduled_end is None:
        return 0
    if row.state in ABSENCE_STATES:
        return 0
    start = datetime.combine(day, row.scheduled_start)
    end = datetime.combine(day, row.scheduled_end)
    if end <= start:
        end += timedelta(days=1)
    seconds = (end - start).total_seconds()
    weekday = day.isoweekday()
    schedule_day = next((d for d in schedule.days.all() if d.weekday == weekday), None)
    if schedule_day is not None:
        for pause in schedule_day.breaks.all():
            if pause.is_paid:
                continue
            a = datetime.combine(day, pause.start_time)
            b = datetime.combine(day, pause.end_time)
            if b <= a:
                b += timedelta(days=1)
            seconds -= (b - a).total_seconds()
    return max(int(seconds), 0)


# --- отчёты -------------------------------------------------------------------


@dataclass
class Table:
    title: str
    columns: list[Column]
    rows: list[dict]
    note: str | None = None


class Report:
    """Общее у всех видов: колонки, область, сводка."""

    main_title = "Данные"
    exact_estimate = True

    def __init__(self, service: BaseService, actor: Actor, spec: ReportSpec,
                 *, author: str) -> None:
        self.service = service
        self.actor = actor
        self.spec = spec
        self.author = author
        self.scope = Scope(service, actor, spec)
        self.columns = columns_for(spec.kind, spec.fields)
        self.notes: list[str] = []

    # переопределяют наследники
    def total_steps(self) -> int:
        raise NotImplementedError

    def rows(self, *, progress: Progress | None = None,
             day_limit: int | None = None) -> Iterator[dict]:
        raise NotImplementedError

    def employee_table(self) -> Table | None:
        return None

    def extra_tables(self) -> list[Table]:
        return []

    def totals(self) -> list[tuple[str, object]]:
        return []

    def employee_count(self) -> int:
        return len(self.scope.assignments())

    # общее
    def summary(self, *, rows: int) -> list[tuple[str, object]]:
        spec = self.spec
        scope = self.scope
        pairs: list[tuple[str, object]] = [
            ("Отчёт", KIND_TITLES[spec.kind]),
            ("Период", f"{spec.first:%d.%m.%Y} — {spec.last:%d.%m.%Y}"),
            ("Офисы", _names([o.name for o in scope.offices], "нет доступных офисов")),
            ("Часовые пояса", _names(sorted({
                f"{o.name}: {zone}" for o, zone in
                ((o, scope.zones[o.id]) for o in scope.offices)
            }), "—")),
            ("Отделы", _names([d.name for d in scope.departments], "все отделы")),
            ("Сотрудник", _full_name(scope.employee) if scope.employee else "все сотрудники"),
            ("Неактивные сотрудники", "включены" if spec.include_inactive else "не включены"),
            ("Сформирован", datetime.now().astimezone().strftime("%d.%m.%Y %H:%M")),
            ("Автор выгрузки", self.author),
            ("Строк в отчёте", rows),
            ("Сотрудников в выборке", self.employee_count()),
        ]
        pairs.extend(self.totals())
        pairs.extend(("Правило", note) for note in self.notes)
        return pairs

    def _identity(self, *, name, number, office, department, day=None) -> dict:
        return {
            "employee": name,
            "employee_number": number,
            "office": office,
            "department": department,
            "date": day,
        }


class _DayReport(Report):
    """Виды по дням: посещаемость, рабочее время, опоздания.

    Объём у них оценочный: «люди × дни» не знает, кто принят или уволен
    посреди периода, поэтому страница пишет «примерно».
    """

    main_title = "По дням"
    exact_estimate = False
    sampled_estimate = False

    def __init__(self, *args, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        from humotech.attendance.hr import AttendanceHrService

        self.hr = AttendanceHrService()
        self.days = self.scope.past_days()
        self.people: dict[uuid.UUID, dict] = {}
        self.schedules = Schedules(self.spec.first, self.spec.last)
        self.notes.append(
            "Сутки считаются в часовом поясе офиса сотрудника; дни после "
            "сегодняшнего в отчёт не попадают."
        )

    def total_steps(self) -> int:
        return len(self.days) * len(self.scope.offices)

    def _presence(self, *, progress, day_limit) -> Iterator[tuple]:
        spec = self.spec
        days = self.days[:day_limit] if day_limit else self.days
        zones = self.scope.zones
        for day in days:
            for office in self.scope.offices:
                if day > today(zones[office.id]):
                    if progress:
                        progress.step()
                    continue
                report = self.hr.presence(
                    self.actor,
                    day=day,
                    office_id=office.id,
                    department_ids=list(spec.department_ids) or None,
                    employee_id=spec.employee_id,
                    include_inactive=spec.include_inactive,
                )
                yield day, office, zones[office.id], report.rows
                if progress:
                    progress.step()

    def _person(self, row) -> dict:
        person = self.people.get(row.employee_id)
        if person is None:
            person = self.people[row.employee_id] = {
                "employee": row.full_name,
                "employee_number": row.employee_number,
                "office": row.office_name,
                "department": row.department_name,
            }
        return person

    def _row_identity(self, row, day) -> dict:
        return self._identity(
            name=row.full_name, number=row.employee_number,
            office=row.office_name, department=row.department_name, day=day,
        )

    def _people_table(self, columns: list[Column]) -> Table:
        base = [Column("employee", "Сотрудник"),
                Column("employee_number", "Табельный номер"),
                Column("office", "Офис"), Column("department", "Отдел")]
        rows = sorted(self.people.values(), key=lambda item: item["employee"] or "")
        return Table("По сотрудникам", base + columns, rows)


class AttendanceReport(_DayReport):
    def __init__(self, *args, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.states: dict[str, int] = {}
        self.event_people: dict[uuid.UUID, set] = {}
        self.notes.append(
            "Выходной и день без графика не считаются отсутствием. Время в "
            "офисе — сумма закрытых сессий; незакрытая сессия отмечена "
            "статусом и во время не входит."
        )

    def rows(self, *, progress=None, day_limit=None):
        for day, office, zone, rows in self._presence(progress=progress, day_limit=day_limit):
            for row in rows:
                status = day_status(row)
                self.states[status] = self.states.get(status, 0) + 1
                person = self._person(row)
                person["present"] = person.get("present", 0) + (1 if row.intervals else 0)
                person["late"] = person.get("late", 0) + (1 if status == "LATE" else 0)
                person["missed"] = person.get("missed", 0) + (1 if status == "NOT_COME" else 0)
                person["absent"] = person.get("absent", 0) + (1 if status in ABSENCE_STATES else 0)
                person["seconds"] = person.get("seconds", 0) + row.seconds
                person["open"] = person.get("open", 0) + sum(
                    1 for part in row.intervals if part.ended_at is None)
                if "marks" in self.spec.fields and row.intervals:
                    self.event_people.setdefault(office.id, set()).add(row.employee_id)
                yield {
                    **self._row_identity(row, day),
                    "schedule": schedule_text(row),
                    "first_entry": local_time(row.first_entry_at, zone),
                    "last_exit": local_time(row.last_exit_at, zone),
                    "marks": marks_text(row.intervals, zone),
                    "office_time": row.seconds,
                    "day_status": status,
                }

    def employee_table(self):
        return self._people_table([
            Column("present", "Дней с отметками", "number"),
            Column("late", "Опозданий", "number"),
            Column("missed", "Не пришёл", "number"),
            Column("absent", "Дней отсутствия", "number"),
            Column("seconds", "Время в офисе", "hours"),
            Column("open", "Незакрытых сессий", "number"),
        ])

    def extra_tables(self):
        if "marks" not in self.spec.fields:
            return []
        from humotech.attendance.models import AttendanceEvent

        columns = [Column("employee", "Сотрудник"),
                   Column("employee_number", "Табельный номер"),
                   Column("office", "Офис"), Column("date", "Дата", "date"),
                   Column("time", "Время", "time"), Column("type", "Событие"),
                   Column("source", "Источник"), Column("verification", "Проверка")]
        rows: list[dict] = []
        zones = self.scope.zones
        offices = {office.id: office for office in self.scope.offices}
        for office_id, people in self.event_people.items():
            zone = zones[office_id]
            start, end = range_bounds(self.spec.first, self.spec.last, zone)
            events = (
                AttendanceEvent.objects.filter(
                    employee_id__in=people, occurred_at__gte=start, occurred_at__lt=end,
                )
                .select_related("employee", "office")
                .order_by("employee__last_name", "employee__first_name", "occurred_at")
            )
            for event in events.iterator(chunk_size=500):
                local = event.occurred_at.astimezone(zone)
                rows.append({
                    "employee": _full_name(event.employee),
                    "employee_number": event.employee.employee_number,
                    "office": event.office.name if event.office else offices[office_id].name,
                    "date": local.date(),
                    "time": local.strftime("%H:%M"),
                    "type": EVENT_TYPES.get(event.event_type, event.event_type),
                    "source": EVENT_SOURCES.get(event.source, event.source),
                    "verification": VERIFICATION.get(
                        event.verification_status, event.verification_status),
                })
                if len(rows) >= XLSX_MAX_ROWS:
                    return [Table("События входа и выхода", columns, rows,
                                  note=f"Показаны первые {XLSX_MAX_ROWS} событий")]
        return [Table("События входа и выхода", columns, rows)]

    def totals(self):
        return [(f"Статус «{DAY_STATUSES[key][0]}»", count)
                for key, count in sorted(self.states.items())]


class WorktimeReport(_DayReport):
    def __init__(self, *args, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.sums = {"planned": 0, "actual": 0, "shortfall": 0, "overtime": 0,
                     "open_sessions": 0}
        self.notes.append(
            "План — смена по графику минус неоплачиваемые перерывы; в "
            "выходной, праздник и дни подтверждённого отсутствия план равен "
            "нулю. Факт — только закрытые сессии. Недостача и переработка "
            "пусты, если графика нет: сравнивать не с чем."
        )

    def rows(self, *, progress=None, day_limit=None):
        for day, office, zone, rows in self._presence(progress=progress, day_limit=day_limit):
            self.schedules.load([row.employee_id for row in rows])
            for row in rows:
                schedule = self.schedules.on(row.employee_id, day)
                planned = planned_seconds(row, schedule, day)
                actual = row.seconds
                if schedule is None:
                    shortfall = overtime = None
                elif row.state in ABSENCE_STATES:
                    shortfall = overtime = 0
                else:
                    shortfall = max(planned - actual, 0)
                    overtime = max(actual - planned, 0)
                closed = sum(1 for part in row.intervals if part.ended_at is not None)
                opened = len(row.intervals) - closed

                person = self._person(row)
                for key, value in (("planned", planned), ("actual", actual),
                                   ("shortfall", shortfall or 0),
                                   ("overtime", overtime or 0),
                                   ("closed_sessions", closed),
                                   ("open_sessions", opened)):
                    person[key] = person.get(key, 0) + value
                    if key in self.sums:
                        self.sums[key] += value

                yield {
                    **self._row_identity(row, day),
                    "planned": planned,
                    "actual": actual,
                    "shortfall": shortfall,
                    "overtime": overtime,
                    "closed_sessions": closed,
                    "open_sessions": opened,
                }

    def employee_table(self):
        return self._people_table([
            Column("planned", "План", "hours"), Column("actual", "Факт", "hours"),
            Column("shortfall", "Недостача", "hours"),
            Column("overtime", "Переработка", "hours"),
            Column("closed_sessions", "Закрытых сессий", "number"),
            Column("open_sessions", "Незакрытых сессий", "number"),
        ])

    def totals(self):
        return [
            ("Плановых часов", _hours(self.sums["planned"])),
            ("Фактических часов", _hours(self.sums["actual"])),
            ("Недостающих часов", _hours(self.sums["shortfall"])),
            ("Переработки, часов", _hours(self.sums["overtime"])),
            ("Незакрытых сессий", self.sums["open_sessions"]),
        ]


class LatenessReport(_DayReport):
    sampled_estimate = True

    def __init__(self, *args, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.count = 0
        self.minutes = 0
        self._reasons: dict[tuple, str] = {}
        self._reasons_loaded: set = set()
        self.notes.append(
            "Опоздание — первый вход позже начала смены по личному графику "
            "сверх допустимого опоздания. Дни без графика и без прихода не "
            "считаются."
        )

    def rows(self, *, progress=None, day_limit=None):
        want_reason = "reason" in self.spec.fields
        for day, office, zone, rows in self._presence(progress=progress, day_limit=day_limit):
            late = [row for row in rows if row.late_minutes]
            if not late:
                continue
            self.schedules.load([row.employee_id for row in late])
            if want_reason:
                self._load_reasons([row.employee_id for row in late], zone)
            for row in late:
                schedule = self.schedules.on(row.employee_id, day)
                self.count += 1
                self.minutes += row.late_minutes
                person = self._person(row)
                person["count"] = person.get("count", 0) + 1
                person["minutes"] = person.get("minutes", 0) + row.late_minutes
                person["worst"] = max(person.get("worst", 0), row.late_minutes)
                yield {
                    **self._row_identity(row, day),
                    "schedule_start": row.scheduled_start.strftime("%H:%M")
                    if row.scheduled_start else None,
                    "first_entry": local_time(row.first_entry_at, zone),
                    "grace": schedule.late_grace_minutes if schedule else None,
                    "late_minutes": row.late_minutes,
                    "reason": self._reasons.get((row.employee_id, day)),
                }

    def _load_reasons(self, employee_ids, zone) -> None:
        from humotech.attendance.models import AttendanceCorrectionRequest

        missing = [key for key in employee_ids if key not in self._reasons_loaded]
        if not missing:
            return
        self._reasons_loaded.update(missing)
        start, end = range_bounds(self.spec.first, self.spec.last, zone)
        requests = (
            AttendanceCorrectionRequest.objects.filter(employee_id__in=missing)
            .filter(
                Q(requested_entry_at__gte=start, requested_entry_at__lt=end)
                | Q(attendance_session__started_at__gte=start,
                    attendance_session__started_at__lt=end)
            )
            .exclude(status="DRAFT")
            .select_related("attendance_session")
            .order_by("submitted_at")
        )
        for request in requests:
            moment = request.requested_entry_at or (
                request.attendance_session.started_at if request.attendance_session else None)
            if moment is None:
                continue
            key = (request.employee_id, moment.astimezone(zone).date())
            title = REQUEST_STATUSES.get(request.status, (request.status,))[0]
            self._reasons.setdefault(
                key, f"Заявка на исправление ({title.lower()}): {request.reason}".strip())

    def employee_table(self):
        return self._people_table([
            Column("count", "Опозданий", "number"),
            Column("minutes", "Минут всего", "minutes"),
            Column("worst", "Наибольшее", "minutes"),
        ])

    def totals(self):
        return [("Опозданий", self.count), ("Минут опозданий", self.minutes)]


class AbsencesReport(Report):
    main_title = "Заявки"

    def __init__(self, *args, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        from humotech.absences.services import AbsenceService

        self.absences = AbsenceService()
        self.people: dict[uuid.UUID, dict] = {}
        self.by_type: dict[str, int] = {}
        self.statuses: dict[str, int] = {}
        self.notes.append(
            "Отобраны заявки, пересекающиеся с периодом. Итоги по дням — "
            "только по одобренным заявкам. Даты — в поясе офиса сотрудника."
        )

    def _queryset(self):
        spec = self.spec
        queryset = self.absences.queue(
            self.actor,
            region_id=spec.region_id,
            date_from=spec.first.isoformat(),
            date_to=spec.last.isoformat(),
        )
        return (
            queryset.filter(employee_id__in=list(self.scope.assignments().keys()))
            .select_related("employee", "absence_type")
            .prefetch_related("documents")
            .order_by("requested_start_at", "id")
        )

    def total_steps(self) -> int:
        return self._queryset().count()

    def rows(self, *, progress=None, day_limit=None):
        assignments = self.scope.assignments()
        zones = self.scope.zones
        for request in self._queryset().iterator(chunk_size=200):
            employee = request.employee
            assignment = assignments.get(employee.id)
            office = assignment.office if assignment else None
            zone = zones.get(office.id) if office else office_zone(None)
            starts = request.requested_start_at.astimezone(zone).date() \
                if request.requested_start_at else None
            ends = request.requested_end_at.astimezone(zone).date() \
                if request.requested_end_at else None
            calendar_days = (ends - starts).days + 1 if starts and ends else None
            working_days = (
                self.absences._working_days(  # noqa: SLF001 — то же правило, что у заявки
                    SimpleNamespace(employee=employee,
                                    organization_id=self.actor.organization_id,
                                    office=office),
                    starts, ends,
                ) if starts and ends and "days" in self.spec.fields else None
            )
            type_name = request.absence_type.name if request.absence_type else None
            self.statuses[request.status] = self.statuses.get(request.status, 0) + 1

            name = _full_name(employee)
            person = self.people.setdefault(employee.id, {
                "employee": name, "employee_number": employee.employee_number,
                "office": office.name if office else None,
                "department": assignment.department.name
                if assignment and assignment.department else None,
            })
            person["requests"] = person.get("requests", 0) + 1
            if request.status == "APPROVED":
                person["approved"] = person.get("approved", 0) + 1
                person["calendar_days"] = person.get("calendar_days", 0) + (calendar_days or 0)
                person["working_days"] = person.get("working_days", 0) + (working_days or 0)
                if type_name:
                    self.by_type[type_name] = self.by_type.get(type_name, 0) + (calendar_days or 0)

            if progress:
                progress.step()
            yield {
                **self._identity(
                    name=name, number=employee.employee_number,
                    office=office.name if office else None,
                    department=assignment.department.name
                    if assignment and assignment.department else None,
                ),
                "absence_type": type_name,
                "starts": starts,
                "ends": ends,
                "calendar_days": calendar_days,
                "working_days": working_days,
                "request_status": request.status,
                "document": document_text(request),
                "decision": decision_text(request, zone),
            }

    def employee_table(self):
        base = [Column("employee", "Сотрудник"), Column("employee_number", "Табельный номер"),
                Column("office", "Офис"), Column("department", "Отдел")]
        rows = sorted(self.people.values(), key=lambda item: item["employee"] or "")
        return Table("По сотрудникам", base + [
            Column("requests", "Заявок", "number"),
            Column("approved", "Одобрено", "number"),
            Column("calendar_days", "Календарных дней (одобрено)", "number"),
            Column("working_days", "Рабочих дней (одобрено)", "number"),
        ], rows)

    def totals(self):
        pairs: list[tuple[str, object]] = [
            (f"Заявок «{REQUEST_STATUSES.get(key, (key,))[0]}»", count)
            for key, count in sorted(self.statuses.items())
        ]
        pairs.extend((f"Дней: {name}", days) for name, days in sorted(self.by_type.items()))
        return pairs


class EmployeesReport(Report):
    main_title = "Сотрудники"

    def __init__(self, *args, **kwargs) -> None:
        super().__init__(*args, **kwargs)
        self.statuses: dict[str, int] = {}
        self.notes.append(
            "Назначение, график и руководитель — на конец периода; для "
            "неактивных — последнее назначение в периоде."
        )

    def total_steps(self) -> int:
        return len(self.scope.assignments())

    def rows(self, *, progress=None, day_limit=None):
        from humotech.telegram.models import TelegramAccount

        assignments = sorted(
            self.scope.assignments().values(),
            key=lambda row: _full_name(row.employee),
        )
        ids = [row.employee_id for row in assignments]
        schedules = Schedules(self.spec.last, self.spec.last)
        if "schedule" in self.spec.fields:
            schedules.load(ids)
        accounts = {}
        if "telegram" in self.spec.fields:
            accounts = {
                row[0]: row[1:] for row in TelegramAccount.objects.filter(
                    employee_id__in=ids).values_list("employee_id", "status", "telegram_username")
            }

        for assignment in assignments:
            employee = assignment.employee
            self.statuses[employee.employment_status] = \
                self.statuses.get(employee.employment_status, 0) + 1
            schedule = schedules.on(employee.id, min(self.spec.last, assignment.valid_to or self.spec.last))
            account = accounts.get(employee.id)
            telegram = None
            if "telegram" in self.spec.fields:
                if account is None:
                    telegram = "Не подключён"
                else:
                    telegram = TELEGRAM.get(account[0], account[0])
                    if account[1]:
                        telegram = f"{telegram} (@{account[1]})"
            if progress:
                progress.step()
            yield {
                "employee": _full_name(employee),
                "employee_number": employee.employee_number,
                "employment_status": employee.employment_status,
                "region": assignment.office.region.name
                if assignment.office and assignment.office.region else None,
                "office": assignment.office.name if assignment.office else None,
                "department": assignment.department.name if assignment.department else None,
                "position": assignment.position.name if assignment.position else None,
                "manager": _full_name(assignment.manager_employee)
                if assignment.manager_employee else None,
                "schedule": schedule.name if schedule else None,
                "hire_date": employee.hire_date,
                "telegram": telegram,
            }

    def totals(self):
        return [(f"Статус «{STATUS_TITLES.get(key, (key,))[0]}»", count)
                for key, count in sorted(self.statuses.items())]


REPORTS = {
    "attendance": AttendanceReport,
    "worktime": WorktimeReport,
    "lateness": LatenessReport,
    "absences": AbsencesReport,
    "employees": EmployeesReport,
}


# --- сервис -------------------------------------------------------------------


class ReportBuilderService(BaseService):
    """Проверка прав, предпросмотр и сборка файла."""

    def require(self, actor: Actor, kind: str) -> None:
        self.access.require(actor, "reports.export")
        self.access.require(actor, KIND_PERMISSIONS[kind])

    def open(self, actor: Actor, spec: ReportSpec, *, author: str) -> Report:
        self.require(actor, spec.kind)
        return REPORTS[spec.kind](self, actor, spec, author=author)

    def check(self, actor: Actor, spec: ReportSpec) -> None:
        """Проверить заказ сразу: права и область, без сборки строк."""
        self.open(actor, spec, author="")

    def preview(self, actor: Actor, spec: ReportSpec, *, fmt: str, limit: int) -> dict:
        report = self.open(actor, spec, author="")
        limit = max(PREVIEW_MIN_ROWS, min(PREVIEW_MAX_ROWS, limit))
        shown: list[dict] = []
        scanned = 0
        sampled_days = None

        if spec.kind in DAY_KINDS:
            sampled_days = min(PREVIEW_DAYS, len(report.days))
            for values in report.rows(day_limit=sampled_days):
                scanned += 1
                if len(shown) < limit:
                    shown.append(values)
                elif not report.sampled_estimate:
                    break
            if report.sampled_estimate:
                estimate = round(scanned / sampled_days * len(report.days)) if sampled_days else 0
            else:
                estimate = report.employee_count() * len(report.days)
        else:
            estimate = report.total_steps()
            for values in report.rows():
                shown.append(values)
                if len(shown) >= limit:
                    break

        warnings: list[str] = []
        if not report.scope.offices:
            warnings.append("Под этими фильтрами нет доступных офисов")
        if spec.kind in DAY_KINDS and len(report.days) < (spec.last - spec.first).days + 1:
            warnings.append("Дни после сегодняшнего в отчёт не попадают")
        if fmt == "xlsx" and estimate > XLSX_MAX_ROWS:
            warnings.append(
                f"В Excel помещается не больше {XLSX_MAX_ROWS:,} строк".replace(",", " ")
                + " — выберите CSV или период короче"
            )

        return {
            "kind": spec.kind,
            "title": KIND_TITLES[spec.kind],
            "date_from": spec.first,
            "date_to": spec.last,
            "fmt": fmt,
            "file_name": spec.file_name(fmt),
            "offices": len(report.scope.offices),
            "employees": report.employee_count(),
            "employee_name": _full_name(report.scope.employee) if report.scope.employee else None,
            "days": len(report.days) if spec.kind in DAY_KINDS else None,
            "rows_estimate": estimate,
            "estimate_exact": report.exact_estimate,
            "sampled_days": sampled_days if getattr(report, "sampled_estimate", False) else None,
            "columns": [{"key": c.key, "title": c.title, "type": c.type}
                        for c in report.columns],
            # На экране — «Фамилия И. О.»: восемь колонок иначе не влезают.
            # В файле ФИО полностью.
            "rows": [[display(_short_name(values.get(c.key))
                              if c.key == "employee" and spec.kind != "employees"
                              else values.get(c.key), c)
                      for c in report.columns]
                     for values in shown],
            "sheets": sheet_titles(report) if fmt == "xlsx" else [],
            "timezones": sorted({str(zone) for zone in report.scope.zones.values()}),
            "warnings": warnings,
        }

    def write(self, actor: Actor, spec: ReportSpec, *, fmt: str, author: str,
              progress: Progress, total: Callable[[int], None] | None = None):
        """Файл отчёта: CSV — кусками строк, XLSX — байтами книги."""
        report = self.open(actor, spec, author=author)
        if total is not None:
            total(report.total_steps())
        if fmt == "csv":
            return csv_chunks(report, progress)
        return xlsx_bytes(report, progress)


def sheet_titles(report: Report) -> list[str]:
    titles = ["Сводка"]
    if report.spec.kind != "employees":
        titles.append("По сотрудникам")
    titles.append(report.main_title)
    if report.spec.kind == "attendance" and "marks" in report.spec.fields:
        titles.append("События входа и выхода")
    return titles


# --- значения -----------------------------------------------------------------


def day_status(row) -> str:
    if row.state in ABSENCE_STATES:
        return row.state
    if row.intervals:
        if row.late_minutes:
            return "LATE"
        if row.open_session_id:
            return "OPEN"
        return "WORKED"
    if row.state in DAY_STATUSES:
        return row.state
    return "NOT_COME"


def schedule_text(row) -> str:
    if row.scheduled_start and row.scheduled_end:
        return f"{row.scheduled_start:%H:%M}–{row.scheduled_end:%H:%M}"
    if row.state == "NO_SCHEDULE":
        return "Нет графика"
    return "Нерабочий день"


def local_time(moment: datetime | None, zone: ZoneInfo) -> str | None:
    return moment.astimezone(zone).strftime("%H:%M") if moment else None


def marks_text(intervals, zone: ZoneInfo) -> str | None:
    if not intervals:
        return None
    parts = []
    for part in intervals:
        start = part.started_at.astimezone(zone).strftime("%H:%M")
        end = part.ended_at.astimezone(zone).strftime("%H:%M") if part.ended_at else "не закрыта"
        parts.append(f"{start}–{end}")
    return "; ".join(parts)


def document_text(request) -> str:
    documents = list(request.documents.all())
    if documents:
        states = {doc.verification_status for doc in documents}
        if "VERIFIED" in states:
            return "Есть, проверен"
        if states == {"REJECTED"}:
            return "Есть, отклонён"
        return "Есть, на проверке"
    if request.absence_type and request.absence_type.requires_document:
        return "Нет, требуется"
    return "Не требуется"


def decision_text(request, zone: ZoneInfo) -> str | None:
    if request.status not in ("APPROVED", "REJECTED") or not request.reviewed_at:
        return None
    title = REQUEST_STATUSES[request.status][0]
    text = f"{title} {request.reviewed_at.astimezone(zone):%d.%m.%Y}"
    if request.review_comment:
        text = f"{text}: {request.review_comment}"
    return text


def display(value, column: Column) -> dict:
    """Клетка предпросмотра: готовый текст и тон для статуса."""
    if value is None:
        return {"text": None, "tone": None}
    kind = column.type
    if kind == "date":
        return {"text": value.strftime("%d.%m.%Y"), "tone": None}
    if kind == "hours":
        minutes = int(value) // 60
        return {"text": f"{minutes // 60} ч {minutes % 60:02d} м", "tone": None}
    if kind == "minutes":
        return {"text": f"{value} мин", "tone": None}
    if kind == "status":
        title, tone = STATUS_TITLES.get(value, (value, "muted"))
        return {"text": title, "tone": tone}
    if kind == "bool":
        return {"text": "да" if value else "нет", "tone": None}
    return {"text": str(value), "tone": None}


def file_value(value, column: Column, *, csv_mode: bool):
    """Значение для файла: числа — числами, статусы — словами, текст — безопасным."""
    if value is None:
        return ""
    kind = column.type
    if kind == "date":
        return value.strftime("%d.%m.%Y") if csv_mode else value
    if kind == "hours":
        hours = round(int(value) / 3600, 2)
        return str(hours).replace(".", ",") if csv_mode else hours
    if kind == "status":
        return STATUS_TITLES.get(value, (value,))[0]
    if kind == "bool":
        return "да" if value else "нет"
    if isinstance(value, str):
        return safe_text(value)
    return value


def csv_chunks(report: Report, progress: Progress) -> Iterator[str]:
    """Одна плоская таблица: заголовок и строки. Без сводки и листов."""
    writer = csv.writer(Echo(), delimiter=";", lineterminator="\r\n")
    yield "﻿"
    yield writer.writerow([column.title for column in report.columns])
    for values in report.rows(progress=progress):
        progress.row()
        yield writer.writerow([
            file_value(values.get(column.key), column, csv_mode=True)
            for column in report.columns
        ])
    progress.finish()


def xlsx_bytes(report: Report, progress: Progress) -> bytes:
    """Книга: «Сводка», «По сотрудникам», данные и события, если нужны."""
    from openpyxl import Workbook

    book = Workbook()
    summary = book.active
    summary.title = "Сводка"
    people = book.create_sheet("По сотрудникам") if report.spec.kind != "employees" else None
    main = book.create_sheet(report.main_title)

    _header(main, report.columns)
    count = 0
    for values in report.rows(progress=progress):
        count += 1
        progress.row()
        if count > XLSX_MAX_ROWS:
            raise ValidationFailed(
                f"В Excel помещается не более {XLSX_MAX_ROWS} строк. "
                "Выберите CSV или период короче.",
                details={"limit": XLSX_MAX_ROWS, "suggested_format": "csv"},
            )
        main.append([file_value(values.get(c.key), c, csv_mode=False)
                     for c in report.columns])
    _finish(main, report.columns)

    if people is not None:
        table = report.employee_table()
        if table is not None:
            _write_table(people, table)
    for table in report.extra_tables():
        _write_table(book.create_sheet(table.title[:31]), table)

    for key, value in report.summary(rows=count):
        summary.append([key, value if isinstance(value, (int, float)) else safe_text(str(value))])
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 80
    from openpyxl.styles import Font

    for cell in summary["A"]:
        cell.font = Font(bold=True)

    progress.finish()
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _header(sheet, columns: list[Column]) -> None:
    from openpyxl.styles import Font, PatternFill

    sheet.append([column.title for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E9E")
    sheet.freeze_panes = "A2"


def _finish(sheet, columns: list[Column]) -> None:
    from openpyxl.utils import get_column_letter

    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = max(12, min(40, len(column.title) + 4))
        if column.type == "date":
            for cell in sheet[letter][1:]:
                cell.number_format = "DD.MM.YYYY"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


def _write_table(sheet, table: Table) -> None:
    _header(sheet, table.columns)
    for values in table.rows:
        sheet.append([file_value(values.get(c.key), c, csv_mode=False) for c in table.columns])
    if table.note:
        sheet.append([])
        sheet.append([table.note])
    _finish(sheet, table.columns)


def _hours(seconds: int) -> float:
    return round(seconds / 3600, 2)


def _names(values: list[str], empty: str) -> str:
    if not values:
        return empty
    if len(values) > 20:
        return ", ".join(values[:20]) + f" и ещё {len(values) - 20}"
    return ", ".join(values)


def _short_name(value):
    """«Абдуллаева Зулфия Нодировна» → «Абдуллаева З. Н.»."""
    if not isinstance(value, str):
        return value
    parts = value.split()
    if len(parts) < 2:
        return value
    return " ".join([parts[0], *(f"{part[0]}." for part in parts[1:3])])


def _full_name(employee) -> str:
    parts = [employee.last_name, employee.first_name, employee.middle_name]
    return " ".join(part for part in parts if part)


__all__ = [
    "BUILDER_VERSION", "PREVIEW_MAX_ROWS", "PREVIEW_MIN_ROWS", "Progress",
    "ReportBuilderService", "ReportSpec", "clean_name", "default_name",
    "is_builder_order",
]
