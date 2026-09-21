"""Конструктор отчётов: поля, предпросмотр, листы, шаблоны и история.

Проверяется то, что обещает страница «Отчёты».

**Предпросмотр — настоящие строки.** Те же построители, что и у файла:
выбранные поля становятся колонками, строки берутся из базы.

**Правила подсчёта.** Выходной — не прогул, будущий день — не «не
пришёл», опоздание — сверх допуска, незакрытая сессия не входит в часы.

**Файл.** Excel — со «Сводкой», «По сотрудникам» и «По дням»; CSV — одна
плоская таблица; имя — вид и период.

**Заказ.** Двойной щелчок не ставит второй отчёт; права проверяются
сразу; скрытая выгрузка пропадает из истории вместе с файлом.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, time, timedelta

import pytest
from django.utils import timezone

from django_tests.conftest import create_actor
from django_tests.test_hr_analytics import FIRST, LAST, five_day_schedule, worked
from humotech.audit.models import AuditLog
from humotech.reports import storage
from humotech.reports.models import ExportJob, ReportTemplate
from humotech.reports.worker import run_once

pytestmark = pytest.mark.django_db

API = "/api/v1"
PERMISSIONS = (
    "reports.export", "employees.read", "attendance.read", "absences.read",
    "offices.read", "schedules.read",
)


@pytest.fixture(autouse=True)
def exports_in_a_temporary_place(settings, tmp_path):
    settings.EXPORTS = {**settings.EXPORTS, "PRIVATE_ROOT": str(tmp_path)}
    return tmp_path


@pytest.fixture()
def hr(api_client, make_user, organization):
    user = make_user(organization, permissions=PERMISSIONS)
    api_client.force_authenticate(user=user)
    api_client.humotech_user = user
    return api_client


def spec(kind="attendance", **extra):
    body = {"kind": kind, "fmt": "xlsx", "date_from": FIRST.isoformat(),
            "date_to": LAST.isoformat()}
    body.update(extra)
    return body


def preview(client, **body):
    return client.post(f"{API}/reports/preview", spec(**body), format="json")


def order(client, **body):
    return client.post(f"{API}/export-jobs/", {**spec(**body), "builder": True},
                       format="json")


def texts(response, column_title):
    data = response.json()
    index = [c["title"] for c in data["columns"]].index(column_title)
    return [row[index]["text"] for row in data["rows"]]


# --- каталог -------------------------------------------------------------------


class TestCatalog:
    def test_five_kinds_with_their_fields(self, hr):
        data = hr.get(f"{API}/reports/catalog").json()
        assert [kind["key"] for kind in data["kinds"]] == [
            "attendance", "worktime", "lateness", "absences", "employees"]
        attendance = data["kinds"][0]
        keys = [item["key"] for item in attendance["fields"]]
        assert "marks" in keys and "day_status" in keys

    def test_the_catalog_needs_the_export_permission(self, api_client, make_user,
                                                     organization):
        user = make_user(organization, permissions=("attendance.read",))
        api_client.force_authenticate(user=user)
        assert api_client.get(f"{API}/reports/catalog").status_code == 403


# --- предпросмотр --------------------------------------------------------------


class TestPreview:
    def test_rows_are_real_and_columns_follow_the_fields(
        self, hr, organization, employee, office
    ):
        five_day_schedule(organization, employee)
        worked(organization, employee, office, FIRST, came=time(9, 0), left=time(18, 9))

        response = preview(hr, fields=["employee", "date", "office_time", "day_status"])
        assert response.status_code == 200, response.content
        data = response.json()
        assert [c["title"] for c in data["columns"]] == ["Сотрудник", "Дата", "Часы", "Статус"]
        # На экране — «Фамилия И. О.», в файле ФИО полностью (см. TestFile).
        assert texts(response, "Сотрудник")[0] == "Иванов И."
        assert texts(response, "Часы")[0] == "9 ч 09 м"
        assert data["rows"][0][3] == {"text": "Рабочий день", "tone": "good"}
        # Один человек, пять дней: оценка — это люди × дни, а не страница.
        assert data["rows_estimate"] == 5
        assert data["offices"] == 1 and data["employees"] == 1
        assert data["file_name"] == "Посещаемость_02.03–06.03.2026.xlsx"
        assert data["sheets"] == ["Сводка", "По сотрудникам", "По дням"]

    def test_a_weekend_is_not_an_absence(self, hr, organization, employee):
        five_day_schedule(organization, employee)
        saturday = LAST + timedelta(days=1)
        response = preview(hr, date_from=saturday.isoformat(), date_to=saturday.isoformat(),
                           fields=["employee", "day_status"])
        assert texts(response, "Статус") == ["Выходной"]

    def test_future_days_are_left_out(self, hr, organization, employee):
        five_day_schedule(organization, employee)
        tomorrow = timezone.localdate() + timedelta(days=1)
        response = preview(hr, date_from=tomorrow.isoformat(),
                           date_to=(tomorrow + timedelta(days=2)).isoformat())
        data = response.json()
        assert data["rows"] == [] and data["rows_estimate"] == 0
        assert "Дни после сегодняшнего в отчёт не попадают" in data["warnings"]

    def test_an_open_session_is_not_counted_as_time(
        self, hr, organization, employee, office
    ):
        five_day_schedule(organization, employee)
        session = worked(organization, employee, office, FIRST)
        session.ended_at = None
        session.duration_seconds = None
        session.exit_event = None
        session.status = "OPEN"
        session.save()

        response = preview(hr, fields=["employee", "office_time", "day_status"])
        assert texts(response, "Часы")[0] == "0 ч 00 м"
        assert texts(response, "Статус")[0] == "Сессия не закрыта"

    def test_lateness_is_measured_over_the_grace(
        self, hr, organization, employee, office
    ):
        five_day_schedule(organization, employee, grace=10)
        worked(organization, employee, office, FIRST, came=time(9, 25))
        worked(organization, employee, office, FIRST + timedelta(days=1), came=time(9, 8))

        response = preview(hr, kind="lateness")
        data = response.json()
        assert len(data["rows"]) == 1
        assert texts(response, "Допуск") == ["10 мин"]
        assert texts(response, "Опоздание") == ["15 мин"]
        assert data["estimate_exact"] is False

    def test_worktime_plans_against_the_schedule(
        self, hr, organization, employee, office
    ):
        five_day_schedule(organization, employee)
        worked(organization, employee, office, FIRST, came=time(9, 0), left=time(17, 0))

        response = preview(hr, kind="worktime", date_to=FIRST.isoformat())
        assert texts(response, "План") == ["9 ч 00 м"]
        assert texts(response, "Факт") == ["8 ч 00 м"]
        assert texts(response, "Недостача") == ["1 ч 00 м"]
        assert texts(response, "Переработка") == ["0 ч 00 м"]

    def test_an_unknown_field_is_refused(self, hr):
        response = preview(hr, fields=["salary"])
        assert response.status_code == 400
        assert "fields" in response.json()["error"]["details"]

    def test_no_fields_is_refused(self, hr):
        assert preview(hr, fields=[]).status_code == 400

    def test_an_office_outside_the_scope_is_refused(
        self, api_client, organization, office, other_office
    ):
        user, _ = create_actor(organization, permissions=PERMISSIONS, office=office)
        api_client.force_authenticate(user=user)
        response = preview(api_client, office_ids=[str(other_office.id)])
        assert response.status_code == 403

    def test_without_filters_only_visible_offices_are_exported(
        self, api_client, organization, employee, office, other_office
    ):
        """Путь, которым идёт страница: офисы не выбраны — берётся область."""
        from humotech.employees.models import Employee, EmployeeAssignment

        stranger = Employee.objects.create(
            organization=organization, employee_number="EMP-0002",
            first_name="Пётр", last_name="Чужой", hire_date=date(2024, 2, 1),
            employment_status="ACTIVE",
        )
        EmployeeAssignment.objects.create(
            organization=organization, employee=stranger, office=other_office,
            employment_type="FULL_TIME", work_mode="ONSITE", is_primary=True,
            valid_from=date(2024, 2, 1),
        )
        user, _ = create_actor(organization, permissions=PERMISSIONS, office=office)
        api_client.force_authenticate(user=user)

        data = preview(api_client, kind="employees").json()
        assert data["offices"] == 1
        assert data["employees"] == 1
        names = [row[0]["text"] for row in data["rows"]]
        assert names == ["Иванов Иван"]

    def test_an_empty_scope_is_explained(self, api_client, organization, employee):
        from humotech.regions.models import Region

        empty = Region.objects.create(organization=organization, code="EMPTY",
                                      name="Пустой", status="ACTIVE")
        user, _ = create_actor(organization, permissions=PERMISSIONS, region=empty)
        api_client.force_authenticate(user=user)

        data = preview(api_client).json()
        assert data["offices"] == 0 and data["rows"] == []
        assert "Под этими фильтрами нет доступных офисов" in data["warnings"]

    def test_the_data_permission_is_required(self, api_client, make_user, organization):
        user = make_user(organization, permissions=("reports.export", "employees.read"))
        api_client.force_authenticate(user=user)
        assert preview(api_client).status_code == 403

    def test_several_offices_and_one_employee(
        self, hr, organization, employee, office, other_office
    ):
        five_day_schedule(organization, employee)
        response = preview(hr, office_ids=[str(office.id), str(other_office.id)],
                           employee_id=str(employee.id), date_to=FIRST.isoformat())
        data = response.json()
        assert data["offices"] == 2
        assert data["employees"] == 1
        assert len(data["rows"]) == 1

    def test_inactive_employees_only_on_request(self, hr, organization, employee):
        employee.employment_status = "TERMINATED"
        employee.save()
        without = preview(hr, kind="employees").json()
        with_them = preview(hr, kind="employees", include_inactive=True).json()
        assert without["employees"] == 0
        assert with_them["employees"] == 1


# --- файл ----------------------------------------------------------------------


class TestFile:
    def test_excel_has_its_sheets_and_name(self, hr, organization, employee, office):
        from openpyxl import load_workbook

        five_day_schedule(organization, employee)
        worked(organization, employee, office, FIRST)
        job_id = order(hr, fields=["employee", "date", "marks", "day_status"]).json()["id"]

        assert run_once() is True
        job = ExportJob.objects.get(id=job_id)
        assert job.status == "SUCCEEDED", job.error_message
        assert job.file_name == "Посещаемость_02.03–06.03.2026.xlsx"
        assert job.progress_total == 5 and job.progress_done == 5
        assert job.total_rows == 5

        with storage.open_export(job.storage_key) as handle:
            book = load_workbook(io.BytesIO(handle.read()))
        assert book.sheetnames == ["Сводка", "По сотрудникам", "По дням",
                                   "События входа и выхода"]
        days = list(book["По дням"].values)
        assert days[0] == ("Сотрудник", "Дата", "Все входы и выходы", "Статус")
        assert days[1][0] == "Иванов Иван"
        events = list(book["События входа и выхода"].values)
        assert [row[5] for row in events[1:]] == ["Вход", "Выход"]

    def test_csv_is_one_flat_table(self, hr, organization, employee):
        five_day_schedule(organization, employee)
        job_id = order(hr, fmt="csv", name="Март: офис/1",
                       fields=["employee", "day_status"]).json()["id"]
        run_once()
        job = ExportJob.objects.get(id=job_id)
        assert job.file_name == "Март офис 1.csv"
        with storage.open_export(job.storage_key) as handle:
            lines = handle.read().decode("utf-8-sig").splitlines()
        assert lines[0] == "Сотрудник;Статус"
        assert len(lines) == 6

    def test_a_double_click_orders_one_report(self, hr):
        key = str(uuid.uuid4())
        first = order(hr, client_request_id=key)
        second = order(hr, client_request_id=key)
        assert first.status_code == 201 and second.status_code == 201
        assert first.json()["id"] == second.json()["id"]
        assert ExportJob.objects.count() == 1

    def test_the_order_is_checked_at_once(self, api_client, make_user, organization):
        user = make_user(organization, permissions=("reports.export",))
        api_client.force_authenticate(user=user)
        assert order(api_client).status_code == 403
        assert ExportJob.objects.count() == 0

    def test_worktime_needs_the_builder(self, hr):
        response = hr.post(f"{API}/export-jobs/", {"kind": "worktime", "fmt": "csv"},
                           format="json")
        assert response.status_code == 400


# --- история -------------------------------------------------------------------


class TestHistory:
    def test_hiding_removes_the_row_and_the_file(self, hr, employee):
        job_id = order(hr, kind="employees").json()["id"]
        run_once()
        job = ExportJob.objects.get(id=job_id)
        key = job.storage_key
        assert storage.exists(key)

        assert hr.post(f"{API}/export-jobs/{job_id}/hide/").status_code == 204
        assert hr.get(f"{API}/export-jobs/").json()["items"] == []
        assert hr.get(f"{API}/export-jobs/counts/").json()["total"] == 0
        assert not storage.exists(key)
        assert AuditLog.objects.filter(action="export.job.hide",
                                       entity_id=job_id).exists()

    def test_a_running_report_is_not_hidden(self, hr):
        job_id = order(hr).json()["id"]
        ExportJob.objects.filter(id=job_id).update(status="RUNNING")
        assert hr.post(f"{API}/export-jobs/{job_id}/hide/").status_code == 409

    def test_an_expired_file_is_counted_apart(self, hr, employee):
        job_id = order(hr, kind="employees").json()["id"]
        run_once()
        ExportJob.objects.filter(id=job_id).update(
            expires_at=timezone.now() - timedelta(minutes=1))

        counts = hr.get(f"{API}/export-jobs/counts/").json()
        assert counts["SUCCEEDED"] == 0 and counts["EXPIRED"] == 1
        row = hr.get(f"{API}/export-jobs/").json()["items"][0]
        assert row["display_status"] == "EXPIRED"
        assert hr.get(f"{API}/export-jobs/", {"status": "SUCCEEDED"}).json()["items"] == []

    def test_parameters_come_back_with_the_row(self, hr, office):
        order(hr, office_ids=[str(office.id)], fields=["employee"], name="Мой")
        row = hr.get(f"{API}/export-jobs/").json()["items"][0]
        assert row["filters"]["office_ids"] == [str(office.id)]
        assert row["filters"]["fields"] == ["employee"]
        assert row["title"] == "Мой"


# --- шаблоны -------------------------------------------------------------------


class TestTemplates:
    def body(self, **extra):
        return {**spec(), "template_name": "Офис за месяц", "period": "this_month",
                "fields": ["employee", "date"], **extra}

    def test_save_list_and_overwrite(self, hr):
        created = hr.post(f"{API}/report-templates/", self.body(), format="json")
        assert created.status_code == 201, created.content
        again = hr.post(f"{API}/report-templates/", self.body(fmt="csv"), format="json")
        assert again.status_code == 200

        items = hr.get(f"{API}/report-templates/").json()["items"]
        assert len(items) == 1
        assert items[0]["fmt"] == "csv"
        assert items[0]["filters"]["period"] == "this_month"
        assert AuditLog.objects.filter(action="report.template.update").exists()

    def test_a_template_is_personal(self, hr, api_client, make_user, organization):
        template_id = hr.post(f"{API}/report-templates/", self.body(),
                              format="json").json()["id"]
        from rest_framework.test import APIClient

        other = APIClient()
        other.force_authenticate(user=make_user(organization, permissions=PERMISSIONS))
        assert other.get(f"{API}/report-templates/").json()["items"] == []
        assert other.delete(f"{API}/report-templates/{template_id}/").status_code == 404

        assert hr.delete(f"{API}/report-templates/{template_id}/").status_code == 204
        assert ReportTemplate.objects.count() == 0

    def test_a_template_needs_a_name(self, hr):
        response = hr.post(f"{API}/report-templates/", self.body(template_name="  "),
                           format="json")
        assert response.status_code == 400


def test_default_names():
    from humotech.reports.builder import default_name

    assert default_name("attendance", date(2026, 8, 15), date(2026, 9, 13)) == \
        "Посещаемость_15.08–13.09.2026"
    assert default_name("lateness", date(2025, 12, 20), date(2026, 1, 5)) == \
        "Опоздания_20.12.2025–05.01.2026"
    assert default_name("employees", date(2026, 8, 1), date(2026, 8, 31)) == \
        "Сотрудники_31.08.2026"
