"""Новые виды выгрузок, период отчёта и счётчики очереди.

Проверяется то, что легко сделать неправильно и незаметно.

**Опоздание — измеренная величина.** День без графика и день без прихода
не превращаются в опоздание: там его не с чем сравнивать. Пустая строка
в отчёте — не «пришёл вовремя».

**Период проверяется при заказе.** Перепутанные местами даты обязаны
стать отказом с указанием поля, а не заданием, которое выглядит принятым
и через минуту умирает.

**Счётчики считают набор, а не страницу.** Число рядом с вкладкой не
должно зависеть от того, какая вкладка открыта.

**Текст из данных остаётся текстом.** Фамилия, начинающаяся со знака
равенства, не должна стать формулой в чужой Excel.
"""

from __future__ import annotations

import io
from datetime import date, time, timedelta

import pytest

from django_tests.test_hr_analytics import FIRST, LAST, five_day_schedule, worked

pytestmark = pytest.mark.django_db

API = "/api/v1"

PERMISSIONS = (
    "reports.export", "employees.read", "attendance.read", "absences.read",
    "analytics.read", "offices.read", "schedules.read",
)


def body_of(response) -> str:
    if hasattr(response, "streaming_content"):
        return b"".join(response.streaming_content).decode("utf-8")
    return response.content.decode("utf-8")


@pytest.fixture(autouse=True)
def exports_in_a_temporary_place(settings, tmp_path):
    settings.EXPORTS = {**settings.EXPORTS, "PRIVATE_ROOT": str(tmp_path)}
    return tmp_path


@pytest.fixture()
def exporter(api_client, make_user, organization):
    user = make_user(organization, permissions=PERMISSIONS)
    api_client.force_authenticate(user=user)
    api_client.humotech_user = user
    return api_client


# --- опоздания ---------------------------------------------------------------


class TestLateness:
    def test_a_late_arrival_is_listed_with_its_schedule(
        self, exporter, organization, employee, office
    ):
        five_day_schedule(organization, employee)
        worked(organization, employee, office, FIRST, came=time(9, 25))

        text = body_of(
            exporter.get(
                f"{API}/reports/lateness/export",
                {"date_from": FIRST.isoformat(), "date_to": LAST.isoformat()},
            )
        )
        assert "Опоздания" in text
        assert "Иванов" in text
        # Именно минуты опоздания, а не время прихода: 09:25 против 09:00.
        assert "25" in text

    def test_a_day_without_a_schedule_is_not_a_late_arrival(
        self, exporter, organization, employee, office
    ):
        """Нет графика — не с чем сравнивать. Это не опоздание.

        Человек пришёл в 11 утра, графика у него нет. Строки в отчёте
        быть не должно: иначе выгрузка обвиняет его в том, чего никто
        не измерял.
        """
        worked(organization, employee, office, FIRST, came=time(11, 0))

        text = body_of(
            exporter.get(
                f"{API}/reports/lateness/export",
                {"date_from": FIRST.isoformat(), "date_to": LAST.isoformat()},
            )
        )
        assert "Иванов" not in text

    def test_an_on_time_arrival_is_not_a_late_arrival(
        self, exporter, organization, employee, office
    ):
        five_day_schedule(organization, employee)
        worked(organization, employee, office, FIRST, came=time(8, 55))

        text = body_of(
            exporter.get(
                f"{API}/reports/lateness/export",
                {"date_from": FIRST.isoformat(), "date_to": LAST.isoformat()},
            )
        )
        assert "Иванов" not in text

    def test_an_empty_row_is_explained_in_the_file(self, exporter, office):
        """Отсутствие строки объяснено в самом файле, а не в переписке."""
        text = body_of(
            exporter.get(
                f"{API}/reports/lateness/export",
                {"date_from": FIRST.isoformat(), "date_to": LAST.isoformat()},
            )
        )
        assert "не означает «пришёл вовремя»" in text


# --- отсутствия --------------------------------------------------------------


class TestAbsences:
    def test_an_absence_lands_in_the_file(
        self, exporter, organization, employee
    ):
        from datetime import datetime, timezone as dt_timezone

        from humotech.absences.models import AbsenceRequest, AbsenceType

        absence_type = AbsenceType.objects.create(
            organization=organization, code="VACATION", name="Отпуск",
            is_paid=True, requires_approval=True,
        )
        AbsenceRequest.objects.create(
            organization=organization,
            employee=employee,
            absence_type=absence_type,
            request_kind="CREATE",
            status="APPROVED",
            requested_start_at=datetime(2026, 3, 3, 0, 0, tzinfo=dt_timezone.utc),
            requested_end_at=datetime(2026, 3, 4, 23, 59, tzinfo=dt_timezone.utc),
            submitted_at=datetime(2026, 2, 20, 8, 0, tzinfo=dt_timezone.utc),
        )

        text = body_of(
            exporter.get(
                f"{API}/reports/absences/export",
                {"date_from": FIRST.isoformat(), "date_to": LAST.isoformat()},
            )
        )
        assert "Иванов" in text
        assert absence_type.name in text
        # Дата подачи — отдельная колонка, и она не равна датам отсутствия.
        assert "2026-02-20" in text

    def test_the_file_says_which_dates_it_filtered_by(self, exporter):
        text = body_of(
            exporter.get(
                f"{API}/reports/absences/export",
                {"date_from": FIRST.isoformat(), "date_to": LAST.isoformat()},
            )
        )
        assert "ПЕРЕСЕКАЮЩИЕСЯ" in text

    def test_reading_absences_needs_its_own_permission(
        self, api_client, make_user, organization
    ):
        """Право на выгрузку не заменяет права на сами данные."""
        user = make_user(
            organization, permissions=("reports.export", "employees.read")
        )
        api_client.force_authenticate(user=user)

        response = api_client.get(
            f"{API}/reports/absences/export",
            {"date_from": FIRST.isoformat(), "date_to": LAST.isoformat()},
        )
        assert response.status_code == 403


# --- посещаемость за период --------------------------------------------------


class TestAttendanceOverAPeriod:
    def test_a_range_gets_a_date_column(
        self, exporter, organization, employee, office
    ):
        worked(organization, employee, office, FIRST)
        worked(organization, employee, office, FIRST + timedelta(days=1))

        text = body_of(
            exporter.get(
                f"{API}/reports/attendance/export",
                {"date_from": FIRST.isoformat(),
                 "date_to": (FIRST + timedelta(days=1)).isoformat()},
            )
        )
        assert "Дата" in text
        assert FIRST.isoformat() in text
        assert (FIRST + timedelta(days=1)).isoformat() in text

    def test_a_single_day_keeps_the_old_shape(
        self, exporter, organization, employee, office
    ):
        """Отчёт за один день не меняется: колонки «Дата» в нём нет."""
        worked(organization, employee, office, FIRST)

        text = body_of(
            exporter.get(
                f"{API}/reports/attendance/export", {"date": FIRST.isoformat()}
            )
        )
        assert "Посещаемость за день" in text
        header = [line for line in text.splitlines() if "Табельный номер" in line]
        assert header and not header[0].startswith("Дата")


# --- период ------------------------------------------------------------------


class TestPeriod:
    def test_reversed_dates_are_refused_at_once(self, exporter):
        """Отказ приходит сразу и указывает поле, а не общий текст."""
        response = exporter.post(
            f"{API}/export-jobs/",
            {"kind": "attendance", "fmt": "csv",
             "date_from": "2026-03-31", "date_to": "2026-03-01"},
            format="json",
        )
        assert response.status_code == 400
        assert "date_to" in response.json()["error"]["details"]

    def test_a_period_longer_than_a_year_is_refused(self, exporter):
        response = exporter.post(
            f"{API}/export-jobs/",
            {"kind": "lateness", "fmt": "csv",
             "date_from": "2020-01-01", "date_to": "2026-01-01"},
            format="json",
        )
        assert response.status_code == 400
        details = response.json()["error"]["details"]
        assert "date_to" in details

    def test_half_a_period_is_refused(self, exporter):
        response = exporter.post(
            f"{API}/export-jobs/",
            {"kind": "sessions", "fmt": "csv", "date_from": "2026-03-01"},
            format="json",
        )
        assert response.status_code == 400
        assert "date_to" in response.json()["error"]["details"]

    def test_the_direct_path_is_guarded_too(self, exporter):
        """Старый путь тоже отказывает, а не держит соединение до таймаута."""
        response = exporter.get(
            f"{API}/reports/lateness/export",
            {"date_from": "2019-01-01", "date_to": "2026-01-01"},
        )
        assert response.status_code == 400


# --- счётчики и фильтры очереди ----------------------------------------------


class TestQueueCounts:
    def make(self, exporter, status: str):
        from humotech.reports.models import ExportJob

        return ExportJob.objects.create(
            organization=exporter.humotech_user.organization,
            requested_by_user=exporter.humotech_user,
            kind="employees",
            fmt="csv",
            status=status,
        )

    def test_counts_cover_the_whole_set(self, exporter):
        for status in ("QUEUED", "RUNNING", "SUCCEEDED", "SUCCEEDED", "FAILED"):
            self.make(exporter, status)

        counts = exporter.get(f"{API}/export-jobs/counts/").json()
        assert counts["total"] == 5
        assert counts["SUCCEEDED"] == 2
        assert counts["CANCELLED"] == 0

    def test_counts_do_not_follow_the_open_tab(self, exporter):
        self.make(exporter, "SUCCEEDED")
        self.make(exporter, "FAILED")

        # Фильтр вкладки в счётчики не передаётся: иначе число рядом
        # с «С ошибкой» менялось бы от того, что открыто.
        counts = exporter.get(
            f"{API}/export-jobs/counts/", {"status": "FAILED"}
        ).json()
        assert counts["total"] == 2

    def test_several_statuses_come_in_one_page(self, exporter):
        self.make(exporter, "QUEUED")
        self.make(exporter, "RUNNING")
        self.make(exporter, "SUCCEEDED")

        page = exporter.get(
            f"{API}/export-jobs/", {"status": "QUEUED,RUNNING"}
        ).json()
        assert len(page["items"]) == 2
        assert {item["status"] for item in page["items"]} == {"QUEUED", "RUNNING"}

    def test_an_unknown_status_is_refused(self, exporter):
        response = exporter.get(f"{API}/export-jobs/", {"status": "ГОТОВО"})
        assert response.status_code == 400

    def test_the_author_comes_with_the_row(self, exporter):
        self.make(exporter, "SUCCEEDED")

        page = exporter.get(f"{API}/export-jobs/").json()
        assert page["items"][0]["requested_by"] == exporter.humotech_user.email


# --- текст из данных ---------------------------------------------------------


class TestFormulaSafety:
    def test_a_name_starting_with_an_equals_sign_stays_text(
        self, exporter, organization, office
    ):
        """Фамилия «=…» не должна стать формулой в чужой Excel."""
        from openpyxl import load_workbook

        from humotech.employees.models import Employee, EmployeeAssignment

        employee = Employee.objects.create(
            organization=organization,
            employee_number="EMP-0777",
            first_name="Иван",
            last_name='=HYPERLINK("http://example.invalid")',
            hire_date=date(2024, 2, 1),
            employment_status="ACTIVE",
        )
        EmployeeAssignment.objects.create(
            organization=organization,
            employee=employee,
            office=office,
            employment_type="FULL_TIME",
            work_mode="ONSITE",
            is_primary=True,
            valid_from=date(2024, 2, 1),
        )

        response = exporter.get(f"{API}/reports/employees/export", {"fmt": "xlsx"})
        book = load_workbook(io.BytesIO(response.content))
        cells = [
            cell
            for row in book.active.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and "HYPERLINK" in cell.value
        ]
        assert cells, "строка сотрудника не попала в файл"
        for cell in cells:
            assert cell.data_type != "f"
            assert cell.value.startswith("'")

    def test_the_same_guard_applies_to_csv(
        self, exporter, organization, office
    ):
        from humotech.employees.models import Employee, EmployeeAssignment

        employee = Employee.objects.create(
            organization=organization,
            employee_number="EMP-0778",
            first_name="Пётр",
            last_name="=1+1",
            hire_date=date(2024, 2, 1),
            employment_status="ACTIVE",
        )
        EmployeeAssignment.objects.create(
            organization=organization,
            employee=employee,
            office=office,
            employment_type="FULL_TIME",
            work_mode="ONSITE",
            is_primary=True,
            valid_from=date(2024, 2, 1),
        )

        text = body_of(exporter.get(f"{API}/reports/employees/export"))
        assert "'=1+1" in text


# --- область видимости после сужения прав ------------------------------------


class TestScopeAfterNarrowing:
    def test_a_narrowed_scope_shrinks_the_next_file(
        self, exporter, organization, office, other_office, employee
    ):
        """Права проверяются при СБОРКЕ: сузили область — файл стал уже.

        Это заявленное свойство модуля, и держится оно на том, что
        исполнитель собирает отчёт от имени заказчика, а не по снимку
        его прав на момент заказа.
        """
        from humotech.accounts.models import UserRoleScope
        from humotech.core.rbac import Actor
        from humotech.employees.models import Employee, EmployeeAssignment
        from humotech.reports.sheets import build_sheet

        far = Employee.objects.create(
            organization=organization,
            employee_number="EMP-0002",
            first_name="Далёкий",
            last_name="Сотрудник",
            hire_date=date(2024, 2, 1),
            employment_status="ACTIVE",
        )
        EmployeeAssignment.objects.create(
            organization=organization,
            employee=far,
            office=other_office,
            employment_type="FULL_TIME",
            work_mode="ONSITE",
            is_primary=True,
            valid_from=date(2024, 2, 1),
        )

        actor = Actor(
            user_id=exporter.humotech_user.id, organization_id=organization.id
        )
        wide = build_sheet("employees", actor, filters={}, author="проверка")
        assert any("Далёкий" in str(row) for row in wide.rows)

        # Область задаётся ЛИБО офисом, ЛИБО регионом: в базе стоит
        # проверка, запрещающая указать оба сразу.
        UserRoleScope.objects.filter(user=exporter.humotech_user).update(
            office=office, region=None
        )

        narrow = build_sheet("employees", actor, filters={}, author="проверка")
        assert not any("Далёкий" in str(row) for row in narrow.rows)

    def test_an_old_file_survives_the_narrowing(
        self, exporter, organization, office, other_office, employee
    ):
        """Уже собранный файл сужением области НЕ отзывается.

        Проверка при скачивании смотрит на владельца, статус и срок
        хранения — но не сравнивает область видимости на момент сборки
        с текущей. Файл, собранный по всей организации, остаётся
        доступным автору и после перевода его в один офис.

        Тест закрепляет фактическое поведение, а не желаемое: пока это
        так, интерфейс не вправе обещать «файлы доступны в пределах
        ваших прав» без оговорки.
        """
        from humotech.accounts.models import UserRoleScope
        from humotech.employees.models import Employee, EmployeeAssignment
        from humotech.reports.models import ExportJob
        from humotech.reports.worker import process_job

        far = Employee.objects.create(
            organization=organization,
            employee_number="EMP-0003",
            first_name="Далёкий",
            last_name="Сотрудник",
            hire_date=date(2024, 2, 1),
            employment_status="ACTIVE",
        )
        EmployeeAssignment.objects.create(
            organization=organization, employee=far, office=other_office,
            employment_type="FULL_TIME", work_mode="ONSITE",
            is_primary=True, valid_from=date(2024, 2, 1),
        )

        job = ExportJob.objects.create(
            organization=organization,
            requested_by_user=exporter.humotech_user,
            kind="employees", fmt="csv", status="RUNNING",
        )
        assert process_job(job) is True

        UserRoleScope.objects.filter(user=exporter.humotech_user).update(
            office=office, region=None
        )

        response = exporter.get(f"{API}/export-jobs/{job.id}/download/")
        assert response.status_code == 200
        text = b"".join(response.streaming_content).decode("utf-8")
        assert "Далёкий" in text


class TestTimezoneInFiles:
    """Файл называет пояс, по которому посчитаны его даты.

    Иначе «01 августа» в отчёте по нескольким офисам означает разное
    для разных читателей, и спорить об этом будет некому.
    """

    def test_attendance_names_the_office_zone(self, exporter, office, employee):
        text = body_of(
            exporter.get(
                f"{API}/reports/attendance/export", {"date": FIRST.isoformat()}
            )
        )
        assert "Часовой пояс" in text

    def test_lateness_names_the_zone(self, exporter, office):
        text = body_of(
            exporter.get(
                f"{API}/reports/lateness/export",
                {"date_from": FIRST.isoformat(), "date_to": LAST.isoformat()},
            )
        )
        assert "Часовой пояс" in text

    def test_sessions_name_the_organization_zone(self, exporter, office, employee):
        """У сессий границы суток отсчитаны по поясу организации.

        Он и назван в файле — раньше строки «Часовой пояс» здесь
        не было вовсе, а экран обещал пояс офиса.
        """
        text = body_of(
            exporter.get(
                f"{API}/reports/sessions/export",
                {"date_from": FIRST.isoformat(), "date_to": LAST.isoformat()},
            )
        )
        assert "Часовой пояс" in text
        assert "по часовому поясу организации" in text
